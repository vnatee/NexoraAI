from __future__ import annotations

import csv
import io
import json
import re
from pathlib import PurePath
from typing import Any
from urllib.parse import quote

import aiohttp
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from core.attachment_extraction import extract_attachment_chunks


MAX_REMOTE_DOCUMENT_BYTES = 25 * 1024 * 1024
OFFICE_MIME = {
    "word": (".docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"),
    "excel": (".xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
    "powerpoint": (".pptx", "application/vnd.openxmlformats-officedocument.presentationml.presentation"),
}
GOOGLE_EXPORTS = {
    "application/vnd.google-apps.document": OFFICE_MIME["word"],
    "application/vnd.google-apps.spreadsheet": OFFICE_MIME["excel"],
    "application/vnd.google-apps.presentation": OFFICE_MIME["powerpoint"],
}


def _safe_name(value: str, extension: str) -> str:
    name = PurePath(str(value or "").replace("\\", "/")).name.strip().strip(".")
    if not name or any(ord(character) < 32 for character in name):
        raise ValueError("Document name is invalid")
    if not name.lower().endswith(extension):
        name += extension
    return name[:240]


def _office_bytes(kind: str, title: str, content: str) -> tuple[str, str, bytes]:
    if kind not in OFFICE_MIME:
        raise ValueError("Document kind must be word, excel, or powerpoint")
    extension, content_type = OFFICE_MIME[kind]
    filename = _safe_name(title, extension)
    output = io.BytesIO()
    if kind == "word":
        document = Document()
        document.add_heading(PurePath(filename).stem, level=1)
        for paragraph in re.split(r"\n\s*\n", content.strip()):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
        document.save(output)
    elif kind == "excel":
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Data"
        rows = list(csv.reader(io.StringIO(content)))
        for row in rows or [[content]]:
            worksheet.append(row)
        workbook.save(output)
    else:
        presentation = Presentation()
        sections = [section.strip() for section in re.split(r"\n\s*---+\s*\n", content) if section.strip()] or [content]
        for section in sections[:100]:
            lines = section.splitlines()
            slide = presentation.slides.add_slide(presentation.slide_layouts[1])
            slide.shapes.title.text = (lines[0] if lines else PurePath(filename).stem)[:200]
            slide.placeholders[1].text = "\n".join(lines[1:])[:10000]
        presentation.save(output)
    return filename, content_type, output.getvalue()


async def _oauth_token(connection: dict[str, Any], secrets: dict[str, str]) -> str:
    provider = connection["provider"]
    config = connection.get("public_config") or {}
    if provider == "microsoft365":
        url = f"https://login.microsoftonline.com/{quote(str(config['tenant_id']))}/oauth2/v2.0/token"
        data = {
            "client_id": config["client_id"],
            "client_secret": secrets.get("client_secret", ""),
            "scope": "https://graph.microsoft.com/.default",
            "grant_type": "client_credentials",
        }
    else:
        url = "https://oauth2.googleapis.com/token"
        data = {
            "client_id": config["client_id"],
            "client_secret": secrets.get("client_secret", ""),
            "refresh_token": secrets.get("refresh_token", ""),
            "grant_type": "refresh_token",
        }
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=30)) as session:
        async with session.post(url, data=data) as response:
            payload = await response.json(content_type=None)
            if response.status >= 400 or not payload.get("access_token"):
                raise ValueError(f"OAuth token request failed ({response.status})")
            return str(payload["access_token"])


async def search_documents(connection: dict[str, Any], secrets: dict[str, str], *, query: str, limit: int = 10) -> list[dict[str, Any]]:
    limit = max(1, min(int(limit), 50))
    token = await _oauth_token(connection, secrets)
    config = connection.get("public_config") or {}
    headers = {"Authorization": f"Bearer {token}"}
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=45)) as session:
        if connection["provider"] == "microsoft365":
            drive_id = quote(str(config["drive_id"]), safe="")
            if query.strip():
                escaped = query[:200].replace("'", "''")
                url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/root/search(q='{quote(escaped)}')"
            else:
                folder = str(config.get("root_folder_id") or "root")
                url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{quote(folder, safe='')}/children" if folder != "root" else f"https://graph.microsoft.com/v1.0/drives/{drive_id}/root/children"
            params = {"$top": str(limit), "$select": "id,name,size,lastModifiedDateTime,webUrl,file,folder,parentReference"}
            async with session.get(url, params=params, headers=headers) as response:
                payload = await response.json(content_type=None)
                if response.status >= 400:
                    raise ValueError(f"Microsoft Graph document search failed ({response.status})")
            return [
                {
                    "id": item.get("id"), "name": item.get("name"), "size": item.get("size"),
                    "modified_at": item.get("lastModifiedDateTime"), "web_url": item.get("webUrl"),
                    "mime_type": (item.get("file") or {}).get("mimeType"), "is_folder": bool(item.get("folder")),
                }
                for item in (payload.get("value") or [])[:limit]
            ]

        terms = [term for term in query[:200].split() if term]
        clauses = ["trashed = false"] + [f"fullText contains '{term.replace(chr(39), '')}'" for term in terms]
        root_folder = str(config.get("root_folder_id") or "").strip()
        if root_folder:
            clauses.append(f"'{root_folder.replace(chr(39), '')}' in parents")
        params: dict[str, str] = {
            "q": " and ".join(clauses),
            "pageSize": str(limit),
            "fields": "files(id,name,mimeType,size,modifiedTime,webViewLink,parents)",
        }
        shared_drive = str(config.get("shared_drive_id") or "").strip()
        if shared_drive:
            params.update({"corpora": "drive", "driveId": shared_drive, "includeItemsFromAllDrives": "true", "supportsAllDrives": "true"})
        async with session.get("https://www.googleapis.com/drive/v3/files", params=params, headers=headers) as response:
            payload = await response.json(content_type=None)
            if response.status >= 400:
                raise ValueError(f"Google Drive document search failed ({response.status})")
        return [
            {
                "id": item.get("id"), "name": item.get("name"), "size": item.get("size"),
                "modified_at": item.get("modifiedTime"), "web_url": item.get("webViewLink"),
                "mime_type": item.get("mimeType"), "is_folder": item.get("mimeType") == "application/vnd.google-apps.folder",
            }
            for item in payload.get("files") or []
        ]


async def read_document(connection: dict[str, Any], secrets: dict[str, str], *, document_id: str) -> dict[str, Any]:
    token = await _oauth_token(connection, secrets)
    headers = {"Authorization": f"Bearer {token}"}
    config = connection.get("public_config") or {}
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=60)) as session:
        if connection["provider"] == "microsoft365":
            drive_id = quote(str(config["drive_id"]), safe="")
            item_id = quote(document_id, safe="")
            async with session.get(f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}", params={"$select": "id,name,size,file,webUrl"}, headers=headers) as response:
                metadata = await response.json(content_type=None)
                if response.status >= 400:
                    raise ValueError(f"Microsoft Graph document lookup failed ({response.status})")
            if int(metadata.get("size") or 0) > MAX_REMOTE_DOCUMENT_BYTES:
                raise ValueError("Document exceeds the 25 MB read limit")
            async with session.get(f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{item_id}/content", headers=headers) as response:
                content = await response.read()
                if response.status >= 400:
                    raise ValueError(f"Microsoft Graph document download failed ({response.status})")
            filename = str(metadata.get("name") or "document")
            content_type = str((metadata.get("file") or {}).get("mimeType") or "application/octet-stream")
            web_url = metadata.get("webUrl")
        else:
            item_id = quote(document_id, safe="")
            params = {"fields": "id,name,mimeType,size,webViewLink"}
            async with session.get(f"https://www.googleapis.com/drive/v3/files/{item_id}", params=params, headers=headers) as response:
                metadata = await response.json(content_type=None)
                if response.status >= 400:
                    raise ValueError(f"Google Drive document lookup failed ({response.status})")
            if int(metadata.get("size") or 0) > MAX_REMOTE_DOCUMENT_BYTES:
                raise ValueError("Document exceeds the 25 MB read limit")
            google_type = str(metadata.get("mimeType") or "")
            if google_type in GOOGLE_EXPORTS:
                extension, content_type = GOOGLE_EXPORTS[google_type]
                filename = _safe_name(str(metadata.get("name") or "document"), extension)
                url = f"https://www.googleapis.com/drive/v3/files/{item_id}/export"
                download_params = {"mimeType": content_type}
            else:
                filename = str(metadata.get("name") or "document")
                content_type = google_type or "application/octet-stream"
                url = f"https://www.googleapis.com/drive/v3/files/{item_id}"
                download_params = {"alt": "media"}
            async with session.get(url, params=download_params, headers=headers) as response:
                content = await response.read()
                if response.status >= 400:
                    raise ValueError(f"Google Drive document download failed ({response.status})")
            web_url = metadata.get("webViewLink")
    if len(content) > MAX_REMOTE_DOCUMENT_BYTES:
        raise ValueError("Document exceeds the 25 MB read limit")
    chunks = extract_attachment_chunks(filename=filename, content_type=content_type, content=content)
    return {
        "id": document_id,
        "name": filename,
        "web_url": web_url,
        "pages": [{"page": item.page_label, "text": item.text} for item in chunks],
    }


async def create_document(connection: dict[str, Any], secrets: dict[str, str], *, kind: str, title: str, content: str) -> dict[str, Any]:
    filename, content_type, payload = _office_bytes(kind, title, content[:200_000])
    token = await _oauth_token(connection, secrets)
    config = connection.get("public_config") or {}
    headers = {"Authorization": f"Bearer {token}"}
    async with aiohttp.ClientSession(timeout=aiohttp.ClientTimeout(total=90)) as session:
        if connection["provider"] == "microsoft365":
            drive_id = quote(str(config["drive_id"]), safe="")
            root_folder_id = str(config.get("root_folder_id") or "").strip()
            if root_folder_id:
                url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/items/{quote(root_folder_id, safe='')}:/{quote(filename, safe='')}:/content"
            else:
                url = f"https://graph.microsoft.com/v1.0/drives/{drive_id}/root:/{quote(filename, safe='')}:/content"
            async with session.put(
                url,
                params={"@microsoft.graph.conflictBehavior": "rename"},
                data=payload,
                headers={**headers, "Content-Type": content_type},
            ) as response:
                result = await response.json(content_type=None)
                if response.status >= 400:
                    raise ValueError(f"Microsoft Graph document create failed ({response.status})")
            return {"created": True, "id": result.get("id"), "name": result.get("name"), "web_url": result.get("webUrl")}

        native_type = {
            "word": "application/vnd.google-apps.document",
            "excel": "application/vnd.google-apps.spreadsheet",
            "powerpoint": "application/vnd.google-apps.presentation",
        }[kind]
        metadata: dict[str, Any] = {"name": PurePath(filename).stem, "mimeType": native_type}
        root_folder = str(config.get("root_folder_id") or "").strip()
        if root_folder:
            metadata["parents"] = [root_folder]
        boundary = "nexora-business-document"
        body = (
            f"--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n{json.dumps(metadata)}\r\n"
            f"--{boundary}\r\nContent-Type: {content_type}\r\n\r\n"
        ).encode() + payload + f"\r\n--{boundary}--\r\n".encode()
        params = {"uploadType": "multipart", "fields": "id,name,mimeType,webViewLink"}
        if config.get("shared_drive_id"):
            params["supportsAllDrives"] = "true"
        async with session.post(
            "https://www.googleapis.com/upload/drive/v3/files",
            params=params,
            data=body,
            headers={**headers, "Content-Type": f"multipart/related; boundary={boundary}"},
        ) as response:
            result = await response.json(content_type=None)
            if response.status >= 400:
                raise ValueError(f"Google Workspace document create failed ({response.status})")
        return {"created": True, "id": result.get("id"), "name": result.get("name"), "web_url": result.get("webViewLink")}


async def test_business_document_connection(connection: dict[str, Any], secrets: dict[str, str]) -> dict[str, Any]:
    items = await search_documents(connection, secrets, query="", limit=1)
    provider = "Microsoft 365" if connection["provider"] == "microsoft365" else "Google Workspace"
    return {"ok": True, "message": f"Connected to {provider}; root access returned {len(items)} item(s)"}
