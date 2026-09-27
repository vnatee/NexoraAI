from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path("docs/questionnaires/NexoraAI_Customer_Discovery_and_Sizing_Questionnaire.xlsx")

NAVY = "172B4D"
BLUE = "1F70C1"
LIGHT_BLUE = "DCEBFA"
PALE_BLUE = "EEF5FC"
TEAL = "008C82"
LIGHT_TEAL = "DDF3F0"
ORANGE = "F59E0B"
LIGHT_ORANGE = "FFF3D6"
RED = "C62828"
LIGHT_RED = "FDE7E7"
GREEN = "2E7D32"
LIGHT_GREEN = "E8F5E9"
GRAY = "667085"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
BLACK = "17202A"
INPUT_FILL = "FFF8D9"

thin_gray = Side(style="thin", color="D0D5DD")


def title(ws, heading: str, subtitle: str, last_col: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = ws.cell(1, 1, heading)
    cell.font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    cell = ws.cell(2, 1, subtitle)
    cell.font = Font(name="Aptos", size=10, color=GRAY, italic=True)
    cell.fill = PatternFill("solid", fgColor=PALE_BLUE)
    cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 32


def style_header(ws, row: int, columns: int) -> None:
    for col in range(1, columns + 1):
        cell = ws.cell(row, col)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=Side(style="medium", color=NAVY))
    ws.row_dimensions[row].height = 32


def style_body(ws, start_row: int, end_row: int, columns: int, input_columns: tuple[int, ...]) -> None:
    for row in range(start_row, end_row + 1):
        ws.row_dimensions[row].height = 44
        for col in range(1, columns + 1):
            cell = ws.cell(row, col)
            cell.font = Font(name="Aptos", size=9, color=BLACK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=thin_gray)
            if col in input_columns:
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
    ws.freeze_panes = f"A{start_row}"
    ws.auto_filter.ref = f"A{start_row - 1}:{get_column_letter(columns)}{end_row}"


def add_table(ws, name: str, start_row: int, end_row: int, columns: int) -> None:
    table = Table(
        displayName=name,
        ref=f"A{start_row}:{get_column_letter(columns)}{end_row}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_list_validation(ws, cell_range: str, formula: str) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True)
    dv.error = "Select a value from the list."
    dv.errorTitle = "Invalid value"
    dv.prompt = "Select the most appropriate value."
    dv.promptTitle = "Customer response"
    dv.showErrorMessage = True
    dv.showInputMessage = True
    ws.add_data_validation(dv)
    dv.add(cell_range)


def question_sheet(wb: Workbook, name: str, heading: str, subtitle: str, rows: list[tuple[str, str, str, str, str, str]]) -> None:
    ws = wb.create_sheet(name)
    title(ws, heading, subtitle, 10)
    headers = [
        "ID",
        "Section / หมวด",
        "Question / คำถาม",
        "Why it matters / เหตุผล",
        "Required",
        "Customer response / คำตอบลูกค้า",
        "Status",
        "Owner",
        "Evidence / Reference",
        "Notes / Open issues",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    style_header(ws, 4, len(headers))
    for row_num, (qid, section, question, why, required, response_type) in enumerate(rows, 5):
        values = [qid, section, question, why, required, "", "Open", "", "", ""]
        for col, value in enumerate(values, 1):
            ws.cell(row_num, col, value)
        if response_type == "yes_no":
            add_list_validation(ws, f"F{row_num}", "=YesNoTBD")
        elif response_type == "deployment":
            add_list_validation(ws, f"F{row_num}", "=DeploymentModels")
        elif response_type == "environment":
            add_list_validation(ws, f"F{row_num}", "=Environments")
        elif response_type == "auth":
            add_list_validation(ws, f"F{row_num}", "=AuthenticationModes")
        elif response_type == "classification":
            add_list_validation(ws, f"F{row_num}", "=DataClassifications")
    end_row = 4 + len(rows)
    style_body(ws, 5, end_row, len(headers), (6, 7, 8, 9, 10))
    add_list_validation(ws, f"G5:G{end_row}", "=QuestionStatus")
    add_table(ws, "Tbl_" + name.replace(" ", "_").replace("&", "and"), 4, end_row, len(headers))
    widths = [13, 23, 58, 42, 12, 42, 16, 20, 34, 36]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def use_case_sheet(wb: Workbook, name: str, heading: str, subtitle: str, cases: list[tuple[str, str, str]]) -> None:
    ws = wb.create_sheet(name)
    title(ws, heading, subtitle, 14)
    headers = [
        "Use case ID",
        "Candidate use case",
        "Example user request",
        "Priority",
        "Personas / Users",
        "Monthly volume",
        "Current process & pain",
        "Knowledge sources",
        "System of record / Integration",
        "Read or Action",
        "Approval required",
        "Data classification",
        "Success metric / Target",
        "Scope decision / Notes",
    ]
    for col, value in enumerate(headers, 1):
        ws.cell(4, col, value)
    style_header(ws, 4, len(headers))
    for row_num, (case_id, case_name, example) in enumerate(cases, 5):
        values = [case_id, case_name, example, "", "", "", "", "", "", "", "", "", "", ""]
        for col, value in enumerate(values, 1):
            ws.cell(row_num, col, value)
    end_row = 4 + len(cases)
    style_body(ws, 5, end_row, len(headers), tuple(range(4, 15)))
    add_list_validation(ws, f"D5:D{end_row}", "=Priorities")
    add_list_validation(ws, f"J5:J{end_row}", "=ReadAction")
    add_list_validation(ws, f"K5:K{end_row}", "=YesNoTBD")
    add_list_validation(ws, f"L5:L{end_row}", "=DataClassifications")
    add_table(ws, "Tbl_" + name.replace(" ", "_"), 4, end_row, len(headers))
    widths = [14, 34, 48, 14, 24, 16, 38, 34, 34, 16, 18, 20, 34, 38]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def build_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "00 Read Me"
    title(
        ws,
        "NexoraAI Customer Discovery & Sizing Questionnaire",
        "แบบสอบถามสำหรับเก็บ Requirement, Configuration และข้อมูล Sizing — Initial scope: HR Employee Services and IT Operations",
        8,
    )
    ws.sheet_view.showGridLines = False
    sections = [
        (4, "Purpose / วัตถุประสงค์", NAVY),
        (5, "Use this workbook before solution design and proposal sizing. Complete customer responses, identify owners/evidence, and close open items during workshops. Yellow cells are customer-input cells.", LIGHT_BLUE),
        (7, "How to use / วิธีใช้", NAVY),
        (8, "1. Complete 01–06 for organization, platform, security, knowledge and integration requirements.\n2. Prioritize HR and IT Operations use cases in 07–08.\n3. Complete all mandatory yellow cells in 09 Sizing Input.\n4. Review 10 Sizing Summary with the solution architect.\n5. Record unresolved decisions and actions in 11 Action Register.", LIGHT_BLUE),
        (10, "Important security instruction", RED),
        (11, "Do not enter passwords, tokens, API keys, client secrets, private keys, connection strings, employee records, or production log samples. Record only the owner and approved secret-store reference.", LIGHT_RED),
        (13, "Sizing disclaimer", ORANGE),
        (14, "All sizing results are budgetary planning baselines—not capacity guarantees. Final production sizing requires a representative pilot/load test using the selected model, context length, document mix, RAG flow, integrations and concurrency profile.", LIGHT_ORANGE),
        (16, "Document control", NAVY),
    ]
    for row, text, color in sections:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=8)
        cell = ws.cell(row, 1, text)
        cell.fill = PatternFill("solid", fgColor=color)
        cell.font = Font(name="Aptos", size=12 if row in (4, 7, 10, 13, 16) else 10, bold=row in (4, 7, 10, 13, 16), color=WHITE if color in (NAVY, RED, ORANGE) else BLACK)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        ws.row_dimensions[row].height = 28 if row in (4, 7, 10, 13, 16) else 64
    controls = [
        (17, "Customer / ลูกค้า", ""),
        (18, "Project / โครงการ", "NexoraAI Discovery"),
        (19, "Prepared by", ""),
        (20, "Customer owner", ""),
        (21, "Workshop date", ""),
        (22, "Workbook version", "1.0"),
        (23, "Generated", date.today().isoformat()),
    ]
    for row, label, value in controls:
        ws.cell(row, 1, label).font = Font(name="Aptos", bold=True, color=NAVY)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        ws.cell(row, 2, value).fill = PatternFill("solid", fgColor=INPUT_FILL if not value else LIGHT_GRAY)
        ws.cell(row, 2).border = Border(bottom=thin_gray)
    for col in range(1, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18

    organization = [
        ("ORG-01", "Organization", "Legal entity and business unit in scope? / นิติบุคคลและหน่วยงานที่อยู่ในขอบเขต", "Defines tenancy, ownership and policy boundary.", "Yes", "text"),
        ("ORG-02", "Stakeholders", "Executive sponsor, business owner, technical owner, security owner and data owners?", "Required for approvals and escalation.", "Yes", "text"),
        ("ORG-03", "Environments", "Which environments are required: Sandbox, Dev, UAT, Production, DR?", "Drives infrastructure, licensing and release process.", "Yes", "environment"),
        ("ORG-04", "Deployment", "Preferred deployment model? / รูปแบบติดตั้งที่ต้องการ", "Selects Azure services, Docker/VM or hybrid architecture.", "Yes", "deployment"),
        ("ORG-05", "Location", "Approved cloud region or on-premises data-center location?", "Data residency, latency and service availability.", "Yes", "text"),
        ("ORG-06", "Network", "Is user access internal-only, VPN, private network, or public HTTPS?", "Defines ingress, DNS, TLS and firewall design.", "Yes", "text"),
        ("ORG-07", "Network", "Is outbound Internet access allowed? If restricted, provide proxy/allowlist process (not credentials).", "Cloud LLM, MCP and SaaS integration connectivity.", "Yes", "text"),
        ("ORG-08", "Availability", "Required service hours, availability target and planned maintenance window?", "Determines topology and operational support.", "Yes", "text"),
        ("ORG-09", "Recovery", "Required RTO and RPO? / เวลากู้คืนและจุดข้อมูลที่ยอมรับได้", "Defines backup, replication and DR design.", "Yes", "text"),
        ("ORG-10", "Growth", "Expected growth over 12 and 36 months for users, documents and requests?", "Prevents sizing only for day-one demand.", "Yes", "text"),
        ("ORG-11", "Support", "Support model, support hours and escalation contacts?", "Required for production operating model.", "No", "text"),
        ("ORG-12", "Compliance", "Applicable standards or regulations (e.g., PDPA, ISO 27001, SEC, internal policy)?", "Defines security and evidence requirements.", "Yes", "text"),
        ("ORG-13", "Procurement", "Target pilot date, production date and procurement constraints?", "Aligns scope and delivery sequence.", "Yes", "text"),
    ]
    question_sheet(wb, "01 Organization", "Organization & Project Context", "ข้อมูลพื้นฐานโครงการ ขอบเขต สถานที่ติดตั้ง Availability และ Timeline", organization)

    platform = [
        ("CFG-01", "URL & TLS", "Production DNS names and certificate source (public CA or internal PKI)?", "Required for HTTPS and IdP callbacks.", "Yes", "text"),
        ("CFG-02", "Authentication", "Authentication methods required? / วิธีเข้าสู่ระบบ", "NexoraAI supports local, Microsoft Entra ID and SAML configuration.", "Yes", "auth"),
        ("CFG-03", "Authentication", "IdP owner, tenant/domain, metadata exchange process and test accounts?", "Required for acceptance testing; do not enter secrets.", "Conditional", "text"),
        ("CFG-04", "Identity", "Which IdP claims identify login, display name, email and groups?", "Required for stable identity mapping.", "Conditional", "text"),
        ("CFG-05", "Access", "List initial Teams, Team owners and approximate member counts.", "Defines workspace and ACL configuration.", "Yes", "text"),
        ("CFG-06", "Access", "Who may be System Admin, Team Admin, Agent Manager, Knowledge Manager and Auditor?", "Supports least privilege and segregation of duties.", "Yes", "text"),
        ("CFG-07", "LLM", "Cloud, local or hybrid LLM strategy? Preferred providers/models?", "Determines connectivity, GPU and quota sizing.", "Yes", "text"),
        ("CFG-08", "LLM", "Required model capabilities: tools/function calling, vision, languages, context length, reasoning?", "Agent and attachment feature compatibility.", "Yes", "text"),
        ("CFG-09", "LLM", "Provider quota: TPM, RPM, concurrent requests and regional capacity?", "Often the primary cloud-LLM bottleneck.", "Conditional", "text"),
        ("CFG-10", "Embedding", "Preferred embedding provider/model and data-residency constraint?", "Controls index compatibility and re-index planning.", "Yes", "text"),
        ("CFG-11", "Document processing", "Native parsing, Azure Document Intelligence or local OCR required?", "Driven by scanned documents, Thai and complex layouts.", "Yes", "text"),
        ("CFG-12", "Storage", "Approved storage: Azure services, local Mongo/MinIO/Qdrant, or hybrid?", "Defines stateful services and backup scope.", "Yes", "text"),
        ("CFG-13", "Retention", "Default chat retention and maximum Team retention: 30/90/180/365 days or no expiry?", "Required for cleanup and compliance policy.", "Yes", "text"),
        ("CFG-14", "License", "Pilot trial or licensed period; expected users/features and offline activation requirement?", "Defines entitlement and activation workflow.", "Yes", "text"),
        ("CFG-15", "Observability", "Preferred monitoring platform and required application/infrastructure metrics?", "Defines system health integration and alerts.", "No", "text"),
    ]
    question_sheet(wb, "02 Platform Config", "NexoraAI Platform Configuration", "Identity, access, LLM, embedding, document processing, storage, retention and licensing", platform)

    security = [
        ("SEC-01", "Data", "Highest data classification permitted in the platform?", "Sets baseline controls and eligible use cases.", "Yes", "classification"),
        ("SEC-02", "Data", "Which PII, payroll, health, performance or regulated fields are prohibited or restricted?", "Required for HR boundaries, guardrails and logging.", "Yes", "text"),
        ("SEC-03", "Encryption", "Encryption requirements at rest and in transit; approved ciphers/certificate policy?", "Defines platform and infrastructure controls.", "Yes", "text"),
        ("SEC-04", "Key management", "Approved key/secret store, rotation interval and key custodians?", "Covers APP_DATA_ENCRYPTION_KEY, provider secrets and certificates.", "Yes", "text"),
        ("SEC-05", "Network", "Required segmentation, private endpoints, firewall, WAF or reverse proxy?", "Defines production network architecture.", "Yes", "text"),
        ("SEC-06", "Guardrails", "Required PII/DLP categories, block/redact/warn behavior and exception approver?", "Creates reusable enforced guardrail profiles.", "Yes", "text"),
        ("SEC-07", "Knowledge ACL", "Required document ACL granularity: personal, Team, Data Group, global?", "NexoraAI enforces retrieval authorization server-side.", "Yes", "text"),
        ("SEC-08", "Tools", "Which integrations/actions are read-only, approval-gated or prohibited?", "Defines tool policy and risk boundary.", "Yes", "text"),
        ("SEC-09", "Audit", "Audit events, viewers, export target and retention period?", "Supports investigation and compliance evidence.", "Yes", "text"),
        ("SEC-10", "Logging", "What content must never appear in logs or telemetry?", "Prevents prompt/tool/secret leakage.", "Yes", "text"),
        ("SEC-11", "Files", "Required malware scanning, file types, maximum size and quarantine process?", "Controls attachment and Knowledge ingestion risk.", "Yes", "text"),
        ("SEC-12", "Vendor risk", "Are cloud LLM data processing, retention and training terms approved?", "Determines eligible providers and deployment mode.", "Conditional", "yes_no"),
        ("SEC-13", "Testing", "Required VA/PT, security review, threat model and acceptance evidence?", "Defines production release gates.", "Yes", "text"),
        ("SEC-14", "Incident", "Security incident owner, reporting SLA and evidence preservation process?", "Required for operating model.", "No", "text"),
    ]
    question_sheet(wb, "03 Security", "Security, Privacy & Governance", "ข้อมูล Classification, PII/DLP, encryption, secrets, ACL, audit และ security acceptance", security)

    knowledge = [
        ("KN-01", "Sources", "List policy, procedure, runbook, architecture, known-error and FAQ repositories.", "Defines initial Knowledge migration scope.", "Yes", "text"),
        ("KN-02", "Formats", "File types and typical content: PDF, Word, Excel, PowerPoint, HTML, images, scans?", "Determines extraction and OCR path.", "Yes", "text"),
        ("KN-03", "Volume", "Current document count, pages, total GB and average file size?", "Primary storage and ingestion sizing input.", "Yes", "text"),
        ("KN-04", "Growth", "New/updated documents per day/month and expected annual growth?", "Worker and storage sizing.", "Yes", "text"),
        ("KN-05", "Languages", "Languages and percentage of Thai/English/multilingual content?", "Affects OCR, embedding and evaluation.", "Yes", "text"),
        ("KN-06", "Scans", "Percentage scanned/image-only and percentage with tables/complex layouts?", "Determines OCR/document intelligence load.", "Yes", "text"),
        ("KN-07", "Ownership", "Document owners, approvers and review/expiry cadence?", "Keeps answers current and accountable.", "Yes", "text"),
        ("KN-08", "ACL", "How are users mapped to document access today?", "Defines Team/Data Group assignment and migration.", "Yes", "text"),
        ("KN-09", "Quality", "Required citation format, page/section evidence and confidence threshold?", "Defines response and acceptance quality.", "Yes", "text"),
        ("KN-10", "Lifecycle", "Archive, deletion, legal hold and source-of-truth synchronization rules?", "Prevents stale or unauthorized retrieval.", "Yes", "text"),
        ("KN-11", "Evaluation", "Provide representative questions, approved answers and negative tests.", "Creates RAG evaluation dataset.", "Yes", "text"),
        ("KN-12", "Freshness", "Maximum acceptable delay from source update to searchable Knowledge?", "Defines ingestion SLA.", "No", "text"),
    ]
    question_sheet(wb, "04 Knowledge & Data", "Knowledge, Documents & Data", "แหล่งข้อมูล ปริมาณ รูปแบบ ภาษา OCR, ACL, freshness และ evaluation dataset", knowledge)

    integrations = [
        ("INT-01", "Inventory", "List required systems and owners: HRIS, Service Desk, CMDB, Datadog/observability, cloud, collaboration, email.", "Creates integration roadmap and ownership.", "Yes", "text"),
        ("INT-02", "Connection", "For each system, is an official hosted MCP, customer-managed MCP, REST API, database or file exchange available?", "Determines connection architecture.", "Yes", "text"),
        ("INT-03", "Environment", "Provide non-production endpoints/tenants and network route owner (no secrets).", "Required for safe integration testing.", "Yes", "text"),
        ("INT-04", "Authentication", "Supported auth: OAuth/OIDC, service principal, PAT/SAT, API key, mTLS? Record owner only.", "Determines credential lifecycle.", "Yes", "text"),
        ("INT-05", "Authorization", "Required scopes/roles and Team/tenant mapping?", "Supports least privilege and tenant isolation.", "Yes", "text"),
        ("INT-06", "Rate limits", "API/MCP rate limits, quotas, result limits and maintenance windows?", "Affects runtime and retry behavior.", "No", "text"),
        ("INT-07", "Actions", "Which operations may change external systems? Who approves each action?", "Separates read-only analysis from governed action.", "Yes", "text"),
        ("INT-08", "Datadog", "Datadog site, tenant count, Team ownership, available toolsets and approved read tools?", "Each tenant should be a separate Team-assigned connection.", "Conditional", "text"),
        ("INT-09", "ITSM", "ServiceNow/Jira projects, ticket fields, status flow and change approval process?", "Required for incident/change workflows.", "Conditional", "text"),
        ("INT-10", "HRIS", "HR system, employee self-service APIs and fields permitted for read/action?", "Extends Knowledge-only HR into transactions safely.", "Conditional", "text"),
        ("INT-11", "Testing", "Integration test data, test accounts and acceptance scenarios?", "Required for end-to-end validation.", "Yes", "text"),
        ("INT-12", "Failure", "Expected behavior when integration is unavailable or returns incomplete data?", "Defines safe fallback and user messaging.", "Yes", "text"),
    ]
    question_sheet(wb, "05 Integrations", "Integrations & MCP Discovery", "ระบบปลายทาง การเชื่อมต่อ สิทธิ์ tenant mapping, rate limit และ approval boundary", integrations)

    operations = [
        ("OPS-01", "Topology", "Single-host pilot or production HA/multi-node requirement?", "Compose is suitable for pilot; HA requires separated tiers.", "Yes", "text"),
        ("OPS-02", "Backup", "Backup destination, frequency, retention, encryption and owner?", "Covers MongoDB, object storage, vectors and configuration.", "Yes", "text"),
        ("OPS-03", "Restore", "Restore test frequency and evidence required?", "Backup is not complete until restore is tested.", "Yes", "text"),
        ("OPS-04", "Monitoring", "CPU, memory, disk, queue, provider quota, worker backlog and certificate alerts?", "System Health is on-demand, not continuous monitoring.", "Yes", "text"),
        ("OPS-05", "Logs", "Central log/SIEM destination and retention?", "Supports operations and audit without sensitive content.", "Yes", "text"),
        ("OPS-06", "Patching", "OS/container/model patch cadence and vulnerability response SLA?", "Defines operational security.", "Yes", "text"),
        ("OPS-07", "Release", "Dev/UAT/Prod promotion, rollback and change approval process?", "Defines deployment governance.", "Yes", "text"),
        ("OPS-08", "Capacity", "Who monitors p95 latency, queue depth, throttling, GPU utilization and ingestion backlog?", "These signals trigger scaling decisions.", "Yes", "text"),
        ("OPS-09", "DR", "DR site/region, failover method and test frequency?", "Required where RTO/RPO cannot be met on one host.", "Conditional", "text"),
        ("OPS-10", "Ownership", "RACI for application, infrastructure, database, vector store, storage, LLM and integrations?", "Avoids support ownership gaps.", "Yes", "text"),
    ]
    question_sheet(wb, "06 Operations", "Deployment & Operating Model", "Topology, backup/restore, monitoring, patching, release, DR และ RACI", operations)

    hr_cases = [
        ("HR-01", "Policy and FAQ assistant", "What is the annual leave carry-forward policy, and when does it expire?"),
        ("HR-02", "Benefits and eligibility guidance", "Am I eligible for this benefit and what conditions apply?"),
        ("HR-03", "Learning request guidance", "What approvals and documents are required before purchasing a course?"),
        ("HR-04", "Employment letter request", "How do I request an English employment letter with salary information?"),
        ("HR-05", "Hybrid/overseas work guidance", "Can I work from another country, and who must approve it?"),
        ("HR-06", "Onboarding assistant", "Give a role-based onboarding checklist for a new employee."),
        ("HR-07", "HR case triage", "Classify my request and route it to the correct HR service queue."),
        ("HR-08", "Employee self-service transaction", "Show my own leave balance or submit a request after confirmation."),
        ("HR-09", "Manager guidance", "What process should a manager follow for a policy exception?"),
        ("HR-10", "Privacy and missing-evidence handling", "Refuse unauthorized employee data and escalate when policy evidence is absent."),
    ]
    use_case_sheet(wb, "07 HR Use Cases", "HR Employee Services — Use Case Discovery", "Prioritize Knowledge-only guidance separately from HRIS actions; never use chat to disclose another employee's confidential data.", hr_cases)

    it_cases = [
        ("IT-01", "Alert and incident triage", "Summarize active critical alerts for production services and rank likely impact."),
        ("IT-02", "Slow application investigation", "Correlate RUM, APM, database, infrastructure, logs, changes and tickets."),
        ("IT-03", "Knowledge/runbook recommendation", "Find the approved runbook and known-error record matching this evidence."),
        ("IT-04", "Evidence timeline", "Build a timestamped incident timeline with source links and confidence."),
        ("IT-05", "Severity trend chart", "Chart alerts by severity for the last three months using exact Datadog results."),
        ("IT-06", "Service Desk enrichment", "Draft or update an incident with evidence, affected service and suggested next steps."),
        ("IT-07", "Change-risk correlation", "Identify recent changes correlated with degradation; do not assert causation without evidence."),
        ("IT-08", "Approval-ready remediation", "Prepare a reversible action plan, validation steps and rollback criteria."),
        ("IT-09", "Cloud inventory and architecture", "Read approved cloud inventory and create an evidence-labelled architecture diagram."),
        ("IT-10", "Post-incident review", "Draft an evidence-backed PIR with gaps and follow-up actions."),
    ]
    use_case_sheet(wb, "08 IT Ops Use Cases", "IT Operations — Use Case Discovery", "Separate observation, recommendation and action. Production-changing tools should be read-only or approval-gated by policy.", it_cases)

    sizing = wb.create_sheet("09 Sizing Input")
    title(sizing, "Sizing Input", "Complete yellow cells. Use measured pilot values where available; otherwise mark the source as Planning assumption.", 8)
    headers = ["ID", "Sizing input", "Unit", "Customer value", "Mandatory", "Source", "Confidence", "Guidance / Notes"]
    for col, value in enumerate(headers, 1):
        sizing.cell(4, col, value)
    style_header(sizing, 4, len(headers))
    metrics = [
        ("SZ-01", "Total named users", "users", "Yes", "Licensing and identity population; not a throughput measure."),
        ("SZ-02", "Monthly active users", "users", "Yes", "Users expected to submit at least one request per month."),
        ("SZ-03", "Peak concurrent users", "users", "Yes", "Users actively using the portal at the same time."),
        ("SZ-04", "Peak simultaneous LLM generations", "generations", "Yes", "Most important runtime concurrency input; measure during pilot."),
        ("SZ-05", "Average AI requests per active user per day", "requests/day", "Yes", "Include Chat, Ask and agent workflows."),
        ("SZ-06", "Peak-hour traffic share", "% of daily", "Yes", "Percentage of daily requests occurring in the busiest hour."),
        ("SZ-07", "Average input/context tokens per request", "tokens", "Yes", "Include prompt, conversation, RAG excerpts and tool context."),
        ("SZ-08", "Average output tokens per request", "tokens", "Yes", "Use measured completion tokens where available."),
        ("SZ-09", "Target p95 end-to-end response time", "seconds", "Yes", "Includes retrieval, tools and model response."),
        ("SZ-10", "Cloud LLM provider TPM quota", "tokens/min", "Conditional", "Enter approved quota; leave blank for local-only."),
        ("SZ-11", "Cloud LLM provider RPM quota", "requests/min", "Conditional", "Enter approved quota; leave blank for local-only."),
        ("SZ-12", "Current Knowledge documents", "documents", "Yes", "Count unique source files."),
        ("SZ-13", "Average source document size", "MB", "Yes", "Use actual repository sampling."),
        ("SZ-14", "Annual Knowledge growth", "%", "Yes", "New and retained source data growth."),
        ("SZ-15", "Average pages per document", "pages", "Yes", "Used for ingestion planning."),
        ("SZ-16", "Scanned/OCR document share", "%", "Yes", "OCR-heavy ingestion consumes more CPU/GPU and time."),
        ("SZ-17", "New or updated documents per business day", "documents/day", "Yes", "Steady-state ingestion volume."),
        ("SZ-18", "Monthly attachment uploads per active user", "files/month", "Yes", "Private Chat/Ask document and image uploads."),
        ("SZ-19", "Average attachment size", "MB", "Yes", "Use actual supported file samples."),
        ("SZ-20", "Attachment/chat retention", "days", "Yes", "Use approved system/team policy."),
        ("SZ-21", "Backup retained copies", "copies", "Yes", "For budgeting; include off-host copies."),
        ("SZ-22", "Required application environments", "count", "Yes", "Dev/UAT/Prod/DR as separate deployments."),
        ("SZ-23", "Production HA required", "Yes/No", "Yes", "Compose single-host is not HA."),
        ("SZ-24", "Deployment model", "model", "Yes", "Cloud LLM, Local LLM, or Hybrid."),
        ("SZ-25", "Local model name and quantization", "text", "Conditional", "Required for local GPU sizing."),
        ("SZ-26", "Measured local model VRAM per replica", "GiB", "Conditional", "Include model weights, KV cache and runtime overhead."),
        ("SZ-27", "Desired local model replicas", "replicas", "Conditional", "Based on concurrency and availability target."),
        ("SZ-28", "Average tool calls per agent request", "calls", "No", "MCP/API/database/search calls increase latency and load."),
        ("SZ-29", "12-month workload growth", "%", "Yes", "Apply to concurrency and requests, not only named users."),
        ("SZ-30", "Required free-disk headroom", "%", "Yes", "Recommended minimum is 25%."),
    ]
    for row_num, (mid, metric, unit, mandatory, guidance) in enumerate(metrics, 5):
        values = [mid, metric, unit, "", mandatory, "", "", guidance]
        for col, value in enumerate(values, 1):
            sizing.cell(row_num, col, value)
    end_sizing = 4 + len(metrics)
    style_body(sizing, 5, end_sizing, len(headers), (4, 6, 7))
    add_list_validation(sizing, f"F5:F{end_sizing}", "=InputSources")
    add_list_validation(sizing, f"G5:G{end_sizing}", "=ConfidenceLevels")
    add_list_validation(sizing, "D27", "=YesNoTBD")
    add_list_validation(sizing, "D28", "=DeploymentModels")
    add_table(sizing, "SizingInputTable", 4, end_sizing, len(headers))
    for col, width in enumerate([13, 46, 18, 22, 14, 22, 16, 54], 1):
        sizing.column_dimensions[get_column_letter(col)].width = width
    sizing.conditional_formatting.add(
        f"D5:D{end_sizing}",
        FormulaRule(formula=["AND($E5=\"Yes\",$D5=\"\")"], fill=PatternFill("solid", fgColor=LIGHT_RED)),
    )

    summary = wb.create_sheet("10 Sizing Summary")
    title(summary, "Indicative Sizing Summary", "Formula-driven planning view. Review with a solution architect and validate by representative load testing before proposal commitment.", 7)
    summary_headers = ["Category", "Metric", "Calculated value", "Unit", "Planning interpretation", "Status", "Owner / Next action"]
    for col, value in enumerate(summary_headers, 1):
        summary.cell(4, col, value)
    style_header(summary, 4, len(summary_headers))
    # Sizing Input customer values are D5:D34. Keep formulas explicit for auditability.
    summary_rows = [
        ("Completeness", "Mandatory inputs completed", '=COUNTIFS(\'09 Sizing Input\'!E5:E34,"Yes",\'09 Sizing Input\'!D5:D34,"<>")', "of mandatory", "Complete required inputs before relying on sizing.", '=IF(C5=COUNTIF(\'09 Sizing Input\'!E5:E34,"Yes"),"Ready","Incomplete")'),
        ("Demand", "Estimated AI requests per business day", "=IFERROR('09 Sizing Input'!D6*'09 Sizing Input'!D9,0)", "requests/day", "Monthly active users × average daily requests.", '=IF(C6>0,"Calculated","Input needed")'),
        ("Demand", "Estimated busiest-hour requests", "=IFERROR(C6*'09 Sizing Input'!D10/100,0)", "requests/hour", "Apply peak-hour share to daily requests.", '=IF(C7>0,"Calculated","Input needed")'),
        ("Demand", "Estimated busiest-hour RPM", "=IFERROR(C7/60,0)", "requests/min", "Compare with provider RPM and measured workflow behavior.", '=IF(C8>0,"Calculated","Input needed")'),
        ("LLM", "Estimated average token demand", "=IFERROR(C8*('09 Sizing Input'!D11+'09 Sizing Input'!D12),0)", "tokens/min", "Planning average during busiest hour; add burst and retry allowance.", '=IF(C9>0,"Calculated","Input needed")'),
        ("LLM", "Recommended quota with 30% headroom", "=ROUNDUP(C9*1.3,0)", "tokens/min", "Compare with approved cloud TPM quota.", '=IF(C10=0,"Input needed",IF(OR(\'09 Sizing Input\'!D14="",\'09 Sizing Input\'!D14=0),"Quota not supplied",IF(C10<=\'09 Sizing Input\'!D14,"Within supplied quota","Quota gap")))'),
        ("Concurrency", "12-month simultaneous generations", "=ROUNDUP(IFERROR('09 Sizing Input'!D8*(1+'09 Sizing Input'!D33/100),0),0)", "generations", "Primary runtime scaling input; validate p95 and queue depth.", '=IF(C11>0,"Calculated","Input needed")'),
        ("Storage", "Current source documents", "=IFERROR('09 Sizing Input'!D16*'09 Sizing Input'!D17/1024,0)", "GiB", "Raw source data only; excludes vectors, metadata and backups.", '=IF(C12>0,"Calculated","Input needed")'),
        ("Storage", "12-month source documents", "=IFERROR(C12*(1+'09 Sizing Input'!D18/100),0)", "GiB", "Raw documents after annual growth.", '=IF(C13>0,"Calculated","Input needed")'),
        ("Storage", "Retained private attachments", "=IFERROR('09 Sizing Input'!D6*'09 Sizing Input'!D22*'09 Sizing Input'!D23*('09 Sizing Input'!D24/30)/1024,0)", "GiB", "Approximation based on monthly uploads and retention.", '=IF(C14>=0,"Calculated","Input needed")'),
        ("Storage", "Primary data planning capacity", "=IFERROR((C13*1.75+C14)*1.25,0)", "GiB", "Includes indicative vector/metadata overhead and 25% operating margin.", '=IF(C15>0,"Budgetary","Input needed")'),
        ("Storage", "Storage including retained backup copies", "=IFERROR(C15*(1+'09 Sizing Input'!D25),0)", "GiB", "Budget across primary and retained copies; placement may be separate.", '=IF(C16>0,"Budgetary","Input needed")'),
        ("Ingestion", "Steady-state pages per business day", "=IFERROR('09 Sizing Input'!D21*'09 Sizing Input'!D19,0)", "pages/day", "Use OCR share to plan worker throughput.", '=IF(C17>0,"Calculated","Input needed")'),
        ("GPU", "Indicative local GPU VRAM incl. 20% margin", "=IFERROR('09 Sizing Input'!D30*'09 Sizing Input'!D31*1.2,0)", "GiB total", "Only valid with measured VRAM per replica; does not select GPU SKU.", '=IF(OR(\'09 Sizing Input\'!D28="Local / On-premises",\'09 Sizing Input\'!D28="Hybrid"),IF(C18>0,"Budgetary","Measurement needed"),"Not required")'),
        ("Infrastructure", "Initial infrastructure tier", '=IF(OR(\'09 Sizing Input\'!D8="",\'09 Sizing Input\'!D8=0),"Input needed",IF(\'09 Sizing Input\'!D8<=10,"Pilot / light cloud-LLM",IF(\'09 Sizing Input\'!D8<=30,"Small production baseline",IF(\'09 Sizing Input\'!D8<=75,"Scale-out production baseline","Capacity engineering required"))))', "tier", "See Recommended Baseline below; not a capacity guarantee.", '=IF(C19="Input needed","Incomplete","Review")'),
    ]
    # The first data row is 5; formula references above intentionally use sheet row numbers.
    for row_num, row in enumerate(summary_rows, 5):
        category, metric, formula, unit, interpretation, status = row
        values = [category, metric, formula, unit, interpretation, status, ""]
        for col, value in enumerate(values, 1):
            summary.cell(row_num, col, value)
    style_body(summary, 5, 19, len(summary_headers), (7,))
    summary["B22"] = "Recommended baseline (budgetary)"
    summary["B22"].font = Font(name="Aptos", size=12, bold=True, color=NAVY)
    baseline_headers = ["Tier", "Starting application baseline", "Data disk", "Topology / Important condition"]
    for col, value in enumerate(baseline_headers, 2):
        summary.cell(23, col, value)
    for col in range(2, 6):
        summary.cell(23, col).fill = PatternFill("solid", fgColor=TEAL)
        summary.cell(23, col).font = Font(name="Aptos", bold=True, color=WHITE)
    baselines = [
        ("Pilot / light cloud-LLM", "8 vCPU / 32 GiB RAM", "256 GiB SSD", "Single host acceptable for evaluation; no HA guarantee."),
        ("Small production baseline", "Start 16 vCPU / 64 GiB RAM", "512 GiB SSD", "Use cloud LLM or separately sized local inference; validate concurrency."),
        ("Scale-out production baseline", "2+ app/worker nodes; size each from load test", "512 GiB+ managed or replicated storage", "Separate app, worker, data and inference tiers; remove single points of failure."),
        ("Capacity engineering required", "Architecture and load-test exercise", "Based on measured working set and retention", "Model, context, tools, OCR, HA and growth materially affect design."),
    ]
    for row_num, values in enumerate(baselines, 24):
        for col, value in enumerate(values, 2):
            summary.cell(row_num, col, value)
            summary.cell(row_num, col).alignment = Alignment(wrap_text=True, vertical="top")
            summary.cell(row_num, col).border = Border(bottom=thin_gray)
        summary.row_dimensions[row_num].height = 46
    for col, width in enumerate([18, 42, 24, 18, 62, 22, 34], 1):
        summary.column_dimensions[get_column_letter(col)].width = width
    summary.freeze_panes = "A5"
    summary.conditional_formatting.add("F5:F19", FormulaRule(formula=['OR(F5="Incomplete",F5="Input needed",F5="Quota gap",F5="Measurement needed")'], fill=PatternFill("solid", fgColor=LIGHT_RED)))
    summary.conditional_formatting.add("F5:F19", FormulaRule(formula=['OR(F5="Ready",F5="Calculated",F5="Within supplied quota")'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)))

    actions = wb.create_sheet("11 Action Register")
    title(actions, "Decision & Action Register", "Track missing information, customer decisions, risks and follow-up evidence discovered during workshops.", 10)
    action_headers = ["Action ID", "Type", "Related question/use case", "Description", "Owner", "Due date", "Status", "Priority", "Decision / Evidence", "Notes"]
    for col, value in enumerate(action_headers, 1):
        actions.cell(4, col, value)
    style_header(actions, 4, len(action_headers))
    for row in range(5, 35):
        actions.cell(row, 1, f"ACT-{row - 4:03d}")
        actions.cell(row, 7, "Open")
    style_body(actions, 5, 34, len(action_headers), tuple(range(2, 11)))
    add_list_validation(actions, "B5:B34", "=ActionTypes")
    add_list_validation(actions, "G5:G34", "=ActionStatus")
    add_list_validation(actions, "H5:H34", "=Priorities")
    add_table(actions, "ActionRegisterTable", 4, 34, len(action_headers))
    for col, width in enumerate([14, 18, 26, 48, 22, 16, 16, 14, 38, 34], 1):
        actions.column_dimensions[get_column_letter(col)].width = width

    lists = wb.create_sheet("_Lists")
    lists.sheet_state = "hidden"
    named_lists = {
        "YesNoTBD": ["Yes", "No", "To be confirmed", "Not applicable"],
        "QuestionStatus": ["Open", "In progress", "Answered", "Evidence pending", "Not applicable"],
        "DeploymentModels": ["Azure managed services", "Azure VM + Docker Compose", "Local / On-premises", "Hybrid", "To be confirmed"],
        "Environments": ["Sandbox only", "Dev + UAT", "Dev + UAT + Production", "Dev + UAT + Production + DR", "To be confirmed"],
        "AuthenticationModes": ["Microsoft Entra ID", "SAML 2.0", "Local authentication", "Entra ID + Local break-glass", "SAML + Local break-glass", "Multiple / To be confirmed"],
        "DataClassifications": ["Public", "Internal", "Confidential", "Restricted / Regulated", "To be confirmed"],
        "Priorities": ["Must have", "Should have", "Could have", "Out of scope", "To be confirmed"],
        "ReadAction": ["Read-only", "Draft only", "Action with approval", "Automatic action", "Not applicable"],
        "InputSources": ["Customer measured", "Customer estimate", "Planning assumption", "Pilot measurement pending"],
        "ConfidenceLevels": ["High", "Medium", "Low", "Unknown"],
        "ActionTypes": ["Information", "Decision", "Risk", "Dependency", "Test", "Security review"],
        "ActionStatus": ["Open", "In progress", "Blocked", "Completed", "Cancelled"],
    }
    from openpyxl.workbook.defined_name import DefinedName

    for col, (list_name, values) in enumerate(named_lists.items(), 1):
        lists.cell(1, col, list_name)
        for row, value in enumerate(values, 2):
            lists.cell(row, col, value)
        column = get_column_letter(col)
        wb.defined_names.add(
            DefinedName(list_name, attr_text=f"'_Lists'!${column}$2:${column}${len(values) + 1}")
        )

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb


def validate(path: Path) -> None:
    workbook = load_workbook(path, data_only=False)
    expected = {
        "00 Read Me",
        "01 Organization",
        "02 Platform Config",
        "03 Security",
        "04 Knowledge & Data",
        "05 Integrations",
        "06 Operations",
        "07 HR Use Cases",
        "08 IT Ops Use Cases",
        "09 Sizing Input",
        "10 Sizing Summary",
        "11 Action Register",
        "_Lists",
    }
    assert expected == set(workbook.sheetnames)
    assert workbook["_Lists"].sheet_state == "hidden"
    assert workbook["10 Sizing Summary"]["C6"].data_type == "f"
    assert workbook["09 Sizing Input"].max_row == 34


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_workbook()
    workbook.save(OUTPUT)
    validate(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
