from __future__ import annotations

from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo


OUTPUT = Path(
    "docs/questionnaires/NexoraAI_Customer_Initial_Questionnaire_Lite.xlsx"
)

NAVY = "172B4D"
BLUE = "1F70C1"
LIGHT_BLUE = "EAF3FC"
TEAL = "008C82"
ORANGE = "F59E0B"
LIGHT_ORANGE = "FFF3D6"
RED = "C62828"
LIGHT_RED = "FDE7E7"
GREEN = "2E7D32"
LIGHT_GREEN = "E8F5E9"
GRAY = "667085"
LIGHT_GRAY = "F2F4F7"
WHITE = "FFFFFF"
BLACK = "17202A"
INPUT_FILL = "FFF8D9"
THIN = Side(style="thin", color="D0D5DD")


def add_title(ws, heading: str, subtitle: str, last_col: int) -> None:
    ws.sheet_view.showGridLines = False
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    ws.cell(1, 1, heading)
    ws.cell(1, 1).font = Font(name="Aptos Display", size=20, bold=True, color=WHITE)
    ws.cell(1, 1).fill = PatternFill("solid", fgColor=NAVY)
    ws.cell(1, 1).alignment = Alignment(vertical="center")
    ws.row_dimensions[1].height = 34
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=last_col)
    ws.cell(2, 1, subtitle)
    ws.cell(2, 1).font = Font(name="Aptos", size=10, color=GRAY, italic=True)
    ws.cell(2, 1).fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[2].height = 30


def add_header(ws, row: int, headers: list[str]) -> None:
    for column, text in enumerate(headers, 1):
        cell = ws.cell(row, column, text)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    ws.row_dimensions[row].height = 32


def style_rows(ws, first: int, last: int, columns: int, answer_columns: tuple[int, ...]) -> None:
    for row in range(first, last + 1):
        ws.row_dimensions[row].height = 48
        for column in range(1, columns + 1):
            cell = ws.cell(row, column)
            cell.font = Font(name="Aptos", size=10, color=BLACK)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=THIN)
            if column in answer_columns:
                cell.fill = PatternFill("solid", fgColor=INPUT_FILL)
    ws.freeze_panes = f"A{first}"


def add_table(ws, name: str, first: int, last: int, columns: int) -> None:
    table = Table(
        displayName=name,
        ref=f"A{first}:{get_column_letter(columns)}{last}",
    )
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(table)


def add_dropdown(ws, cell_range: str, values: list[str]) -> None:
    escaped = ",".join(values).replace('"', '""')
    validation = DataValidation(
        type="list", formula1=f'"{escaped}"', allow_blank=True
    )
    validation.error = "Please select a value from the list."
    validation.errorTitle = "Invalid value"
    validation.showErrorMessage = True
    ws.add_data_validation(validation)
    validation.add(cell_range)


def question_sheet(
    wb: Workbook,
    name: str,
    heading: str,
    subtitle: str,
    questions: list[tuple[str, str, str]],
) -> None:
    ws = wb.create_sheet(name)
    add_title(ws, heading, subtitle, 6)
    headers = [
        "No.",
        "Topic / หัวข้อ",
        "Question / คำถาม",
        "Example or hint / ตัวอย่าง",
        "Customer answer / คำตอบลูกค้า",
        "Owner / Notes",
    ]
    add_header(ws, 4, headers)
    for row_number, (topic, question, hint) in enumerate(questions, 5):
        values = [row_number - 4, topic, question, hint, "", ""]
        for column, value in enumerate(values, 1):
            ws.cell(row_number, column, value)
    last = 4 + len(questions)
    style_rows(ws, 5, last, 6, (5, 6))
    add_table(ws, f"Tbl_{name.replace(' ', '_').replace('&', 'and')}", 4, last, 6)
    for column, width in enumerate([8, 22, 52, 44, 44, 30], 1):
        ws.column_dimensions[get_column_letter(column)].width = width


def use_case_sheet(
    wb: Workbook,
    name: str,
    heading: str,
    subtitle: str,
    cases: list[tuple[str, str]],
) -> None:
    ws = wb.create_sheet(name)
    add_title(ws, heading, subtitle, 8)
    headers = [
        "No.",
        "Use case / งานที่ต้องการ",
        "Example question / ตัวอย่างคำถาม",
        "Interested?",
        "Priority",
        "Current problem / ปัญหาปัจจุบัน",
        "Information or system used today",
        "Expected result / ผลลัพธ์ที่ต้องการ",
    ]
    add_header(ws, 4, headers)
    for row_number, (case, example) in enumerate(cases, 5):
        values = [row_number - 4, case, example, "", "", "", "", ""]
        for column, value in enumerate(values, 1):
            ws.cell(row_number, column, value)
    last = 4 + len(cases)
    style_rows(ws, 5, last, 8, (4, 5, 6, 7, 8))
    add_dropdown(ws, f"D5:D{last}", ["Yes", "No", "Maybe"])
    add_dropdown(
        ws,
        f"E5:E{last}",
        ["High / สูง", "Medium / กลาง", "Low / ต่ำ", "Later / ภายหลัง"],
    )
    add_table(ws, f"Tbl_{name.replace(' ', '_')}", 4, last, 8)
    for column, width in enumerate([8, 34, 50, 14, 17, 38, 38, 38], 1):
        ws.column_dimensions[get_column_letter(column)].width = width


def build_workbook() -> Workbook:
    wb = Workbook()
    cover = wb.active
    cover.title = "00 Start Here"
    add_title(
        cover,
        "NexoraAI Initial Customer Questionnaire — Lite",
        "แบบสอบถามข้อมูลเบื้องต้นสำหรับ HR Employee Services และ IT Operations | ใช้เวลาประมาณ 30–45 นาที",
        7,
    )
    cover.sheet_view.showGridLines = False
    blocks = [
        (4, "Purpose / วัตถุประสงค์", NAVY, WHITE),
        (5, "ใช้สำหรับพูดคุยครั้งแรก เพื่อเข้าใจเป้าหมาย Use Case ข้อมูล ระบบที่เกี่ยวข้อง จำนวนผู้ใช้ และรูปแบบติดตั้ง ก่อนจัดทำ Demo, Pilot และงบประมาณเบื้องต้น", LIGHT_BLUE, BLACK),
        (7, "How to complete / วิธีกรอก", NAVY, WHITE),
        (8, "กรอกเฉพาะช่องสีเหลือง หากยังไม่ทราบให้ใส่ “TBC” ได้ เลือก Use Case ที่สนใจและให้ Priority จากนั้นกรอก Sizing เท่าที่ทราบ ตัวเลขทั้งหมดสามารถยืนยันเพิ่มเติมใน Technical Workshop", LIGHT_BLUE, BLACK),
        (10, "Security note / ข้อควรระวัง", RED, WHITE),
        (11, "ห้ามใส่ Password, API Key, Token, Client Secret, Connection String หรือข้อมูลพนักงานจริงในไฟล์นี้ ให้ระบุเพียงชื่อระบบและผู้ดูแล", LIGHT_RED, BLACK),
        (13, "Sizing note", ORANGE, WHITE),
        (14, "Sizing ในไฟล์นี้เป็นเพียงแนวทางเบื้องต้นสำหรับเสนอ Pilot/Budget ต้องทดสอบด้วย Model, เอกสาร, Integration และจำนวนผู้ใช้จริงก่อน Production", LIGHT_ORANGE, BLACK),
    ]
    for row, text, fill, color in blocks:
        cover.merge_cells(start_row=row, start_column=1, end_row=row, end_column=7)
        cell = cover.cell(row, 1, text)
        cell.fill = PatternFill("solid", fgColor=fill)
        cell.font = Font(name="Aptos", size=12 if row in (4, 7, 10, 13) else 10, bold=row in (4, 7, 10, 13), color=color)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cover.row_dimensions[row].height = 30 if row in (4, 7, 10, 13) else 54
    cover.cell(17, 1, "Customer / ลูกค้า").font = Font(bold=True, color=NAVY)
    cover.cell(18, 1, "Contact person").font = Font(bold=True, color=NAVY)
    cover.cell(19, 1, "Workshop date").font = Font(bold=True, color=NAVY)
    cover.cell(20, 1, "Prepared by").font = Font(bold=True, color=NAVY)
    for row in range(17, 21):
        cover.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
        cover.cell(row, 2).fill = PatternFill("solid", fgColor=INPUT_FILL)
        cover.cell(row, 2).border = Border(bottom=THIN)
    cover.cell(22, 1, "Version")
    cover.cell(22, 2, "Lite 1.0")
    cover.cell(23, 1, "Generated")
    cover.cell(23, 2, date.today().isoformat())
    for column in range(1, 8):
        cover.column_dimensions[get_column_letter(column)].width = 20

    basic_questions = [
        ("Business goal", "What problem do you want NexoraAI to solve first? / อยากให้ระบบช่วยแก้ปัญหาอะไรเป็นอันดับแรก", "เช่น ลดเวลาตอบคำถาม HR หรือลดเวลาวิเคราะห์ Incident"),
        ("Scope", "Which department or business unit will join the pilot?", "HR, IT Operations, Service Desk หรือหน่วยงานที่เลือก"),
        ("Users", "Who will use the system? / ผู้ใช้งานหลักคือใคร", "พนักงานทั่วไป, HR, Service Desk, NOC, System Engineer, Manager"),
        ("Timeline", "When would you like to start the pilot and production use?", "เดือน/ไตรมาสที่ต้องการ"),
        ("Deployment", "Where should the system run?", "Azure, On-premises, Hybrid หรือยังไม่ตัดสินใจ"),
        ("AI model", "Can the system use cloud AI, or must AI run inside the organization?", "Cloud AI, Local AI, Hybrid, TBC"),
        ("Login", "How should users sign in? / ต้องการ Login แบบใด", "Microsoft Entra ID, SAML, Local user"),
        ("Access", "Will different departments need separate Teams and separate data access?", "เช่น HR เห็นเฉพาะ HR, IT เห็นเฉพาะ IT"),
        ("Language", "Which languages must the system support?", "Thai, English หรือทั้งสองภาษา"),
        ("Success", "How will you know the pilot is successful?", "เช่น ลดเวลาค้นหา 50%, คำตอบถูกต้อง 90%, ลด MTTR"),
    ]
    question_sheet(wb, "01 Basic Info", "Customer & Project Overview", "คำถามพื้นฐานเพื่อกำหนดเป้าหมาย ขอบเขต และรูปแบบโครงการ", basic_questions)

    hr_cases = [
        ("Answer HR policy questions", "วันลาคงเหลือใช้ได้ถึงเมื่อไร และมีเงื่อนไขอะไร"),
        ("Explain benefits and eligibility", "ผมมีสิทธิ์เบิกสวัสดิการนี้หรือไม่"),
        ("Guide employee requests", "ต้องเตรียมอะไรและขออนุมัติใคร"),
        ("Employment letter guidance", "ขอหนังสือรับรองการทำงานอย่างไร ใช้เวลากี่วัน"),
        ("Learning and training guidance", "หลักสูตรนี้เบิกได้หรือไม่ ต้องอนุมัติอย่างไร"),
        ("Onboarding assistant", "พนักงานใหม่ต้องทำอะไรบ้างในสัปดาห์แรก"),
        ("Route HR cases", "ช่วยส่งคำขอนี้ไปยังทีม HR ที่ถูกต้อง"),
        ("Employee self-service action", "ดูข้อมูลของตนเองหรือส่งคำขอหลังยืนยัน"),
        ("Privacy-safe response", "ไม่เปิดเผยเงินเดือนหรือผลประเมินของผู้อื่น"),
    ]
    use_case_sheet(wb, "02 HR Use Cases", "HR Employee Services", "เลือกงานที่สนใจสำหรับ Demo/Pilot ก่อน เริ่มจาก Knowledge และเพิ่ม HR system action ภายหลังได้", hr_cases)

    it_cases = [
        ("Summarize alerts and incidents", "ตอนนี้มี Critical Alert อะไร และกระทบบริการใด"),
        ("Investigate a slow application", "ทำไม Trading Application ช้าในช่วง 30 นาทีที่ผ่านมา"),
        ("Correlate monitoring evidence", "เชื่อมโยง APM, Logs, Database, Infrastructure และ Change"),
        ("Find approved runbooks", "มี Runbook หรือ Known Error ใดตรงกับอาการนี้"),
        ("Create an evidence timeline", "เรียงเหตุการณ์และหลักฐานตามเวลา"),
        ("Create charts and trends", "ทำกราฟ Alert แยกตาม Severity ย้อนหลัง 3 เดือน"),
        ("Draft a Service Desk ticket", "สรุปข้อมูลเพื่อเปิด Incident ใน ServiceNow/Jira"),
        ("Recommend remediation", "เสนอวิธีแก้ พร้อมความเสี่ยง ขั้นตอนตรวจสอบ และ Rollback"),
        ("Draft post-incident review", "สร้าง PIR จากหลักฐานที่เกิดขึ้นจริง"),
    ]
    use_case_sheet(wb, "03 IT Ops Use Cases", "IT Operations", "เลือกงานที่ต้องการเริ่มก่อน และระบุว่าต้องการเพียงอ่าน/แนะนำ หรือให้ดำเนินการหลังอนุมัติ", it_cases)

    data_questions = [
        ("HR knowledge", "Where are HR policies and procedures stored today?", "SharePoint, File server, Website, PDF/Word, Google Drive"),
        ("IT knowledge", "Where are runbooks, known errors and architecture documents stored?", "SharePoint, Confluence, File server, ServiceNow KB"),
        ("Document size", "Approximately how many documents and how much total storage?", "เช่น 5,000 ไฟล์ / 50 GB; ประมาณการได้"),
        ("Document type", "Are documents mostly text files or scanned images?", "PDF/Word/Excel/PPT, รูปสแกน, ภาษาไทย/อังกฤษ"),
        ("Update", "How often are documents added or changed?", "Daily, weekly, monthly"),
        ("Permissions", "Do some documents have restricted access?", "เช่น Payroll เฉพาะ HR, Runbook เฉพาะ IT"),
        ("HR system", "Which HR system is used, and is API access available?", "ระบุชื่อระบบและผู้ดูแล ไม่ต้องใส่ credential"),
        ("Monitoring", "Which monitoring tools are used?", "Datadog, Azure Monitor, Grafana, other"),
        ("Service Desk", "Which ticket/change system is used?", "ServiceNow, Jira Service Management, other"),
        ("Cloud", "Which cloud platforms should the Agent read?", "Azure, AWS, GCP, On-premises"),
        ("Connection", "Can NexoraAI connect to these systems through HTTPS/API/MCP?", "ระบุข้อจำกัด Firewall, Proxy หรือ Private Network"),
        ("Actions", "Should Agents only read information or also create/update records?", "เริ่มจาก Read-only หรือ Action with approval"),
    ]
    question_sheet(wb, "04 Data & Systems", "Data, Knowledge & Connected Systems", "ข้อมูลเบื้องต้นเกี่ยวกับเอกสาร ระบบต้นทาง และวิธีเชื่อมต่อ", data_questions)

    security_questions = [
        ("Sensitive data", "What sensitive information may be present?", "PII, เงินเดือน, สุขภาพ, ผลประเมิน, ข้อมูลลูกค้า, Log production"),
        ("Restrictions", "What information must the AI never show or use?", "ระบุข้อมูลต้องห้ามหรือเฉพาะบทบาท"),
        ("Approval", "Which AI actions require human approval?", "เปิด Ticket, เปลี่ยนระบบ, ส่ง Email, ส่ง HR request"),
        ("Retention", "How long should chats and uploaded files be kept?", "30, 90, 180, 365 วัน หรือระบุเอง"),
        ("Audit", "Who needs to review usage and audit records?", "Security, Compliance, Internal Audit, System Admin"),
        ("Availability", "When must the service be available?", "Business hours, 24x7, planned maintenance window"),
        ("Recovery", "How quickly must service and data be recovered?", "ระบุ RTO/RPO หากทราบ หรือ TBC"),
        ("Compliance", "Are there specific policies or regulations to follow?", "PDPA, ISO 27001, SEC, internal policy"),
    ]
    question_sheet(wb, "05 Security & Ops", "Security & Operations — Initial Questions", "คำถามระดับต้น รายละเอียดเชิงเทคนิคสามารถทำต่อใน Security/Architecture Workshop", security_questions)

    sizing = wb.create_sheet("06 Simple Sizing")
    add_title(sizing, "Simple Sizing Input", "กรอกเท่าที่ทราบ ใช้เพื่อประเมิน Pilot และงบประมาณเบื้องต้นเท่านั้น", 7)
    headers = ["No.", "Sizing question / คำถาม", "Unit", "Customer answer", "Example / Hint", "Source", "Notes"]
    add_header(sizing, 4, headers)
    sizing_rows = [
        ("How many total users may have access?", "users", "", "เช่น 500"),
        ("How many users may use it in one day?", "users/day", "", "เช่น 100"),
        ("How many users may use it at the same time?", "concurrent users", "", "เช่น 20"),
        ("About how many AI questions per user per day?", "questions/day", "", "เช่น 5"),
        ("How many Knowledge documents are in scope?", "documents", "", "เช่น 5,000"),
        ("Approximate total document storage?", "GB", "", "เช่น 50 GB"),
        ("How many new/updated documents per month?", "documents/month", "", "เช่น 200"),
        ("What percentage of documents are scanned images?", "%", "", "เช่น 20%"),
        ("How long should chats/files be retained?", "days", "", "เช่น 90 วัน"),
        ("How many environments are needed?", "environments", "", "Pilot, UAT, Production, DR"),
        ("Is production High Availability required?", "Yes/No", "", "หากยังไม่ทราบเลือก TBC"),
        ("Preferred AI deployment", "model", "", "Cloud AI, Local AI, Hybrid, TBC"),
        ("Expected growth in the next 12 months?", "%", "", "จำนวนผู้ใช้/คำถาม/เอกสาร"),
    ]
    for row_number, (question, unit, answer, hint) in enumerate(sizing_rows, 5):
        values = [row_number - 4, question, unit, answer, hint, "", ""]
        for column, value in enumerate(values, 1):
            sizing.cell(row_number, column, value)
    last = 4 + len(sizing_rows)
    style_rows(sizing, 5, last, 7, (4, 6, 7))
    add_dropdown(sizing, f"F5:F{last}", ["Customer confirmed", "Customer estimate", "Initial assumption", "TBC"])
    add_dropdown(sizing, "D15", ["Yes", "No", "TBC"])
    add_dropdown(sizing, "D16", ["Cloud AI", "Local AI", "Hybrid", "TBC"])
    add_table(sizing, "Tbl_Simple_Sizing", 4, last, 7)
    for column, width in enumerate([8, 52, 20, 25, 38, 22, 32], 1):
        sizing.column_dimensions[get_column_letter(column)].width = width

    result = wb.create_sheet("07 Initial Summary")
    add_title(result, "Initial Planning Summary", "สรุปอัตโนมัติระดับต้น หลังกรอก Simple Sizing — ต้องยืนยันอีกครั้งด้วย Technical Workshop และ Load Test", 6)
    summary_headers = ["Item", "Result", "Meaning", "Customer decision", "Owner", "Next step"]
    add_header(result, 4, summary_headers)
    summary_rows = [
        ("Input completeness", '=COUNTIF(\'06 Simple Sizing\'!D5:D17,"<>")&" / 13 answered"', "Complete missing answers before proposal sizing.", "", "", ""),
        ("Daily question estimate", "=IFERROR('06 Simple Sizing'!D6*'06 Simple Sizing'!D8,0)", "Daily users × questions per user.", "", "", ""),
        ("Initial workload class", '=IF(OR(\'06 Simple Sizing\'!D7="",\'06 Simple Sizing\'!D7=0),"Input needed",IF(\'06 Simple Sizing\'!D7<=10,"Pilot / Light",IF(\'06 Simple Sizing\'!D7<=30,"Small Production",IF(\'06 Simple Sizing\'!D7<=75,"Scale-out Review","Capacity Workshop Required"))))', "Based mainly on concurrent users; AI model and workflow may change the result.", "", "", ""),
        ("Starting infrastructure guidance", '=IF(B7="Pilot / Light","8 vCPU / 32 GB RAM / 256 GB SSD",IF(B7="Small Production","Start 16 vCPU / 64 GB RAM / 512 GB SSD",IF(B7="Scale-out Review","Plan separate/replicated application and data tiers","Technical sizing required")))', "Budgetary baseline only. Local AI requires separate GPU sizing.", "", "", ""),
        ("High Availability", "='06 Simple Sizing'!D15", "Single-host Docker Compose is not HA.", "", "", ""),
        ("AI deployment", "='06 Simple Sizing'!D16", "Local AI requires model/GPU/context/concurrency assessment.", "", "", ""),
        ("Knowledge volume", "='06 Simple Sizing'!D9", "Use actual document sample to validate extraction and RAG quality.", "", "", ""),
        ("Next gate", '=IF(COUNTIF(\'06 Simple Sizing\'!D5:D17,"<>")<10,"Complete discovery","Define pilot scope and success criteria")', "Proceed to detailed technical/security questionnaire only after initial scope agreement.", "", "", ""),
    ]
    for row_number, values in enumerate(summary_rows, 5):
        for column, value in enumerate(values, 1):
            result.cell(row_number, column, value)
    style_rows(result, 5, 12, 6, (4, 5, 6))
    add_table(result, "Tbl_Initial_Summary", 4, 12, 6)
    for column, width in enumerate([30, 34, 56, 34, 22, 40], 1):
        result.column_dimensions[get_column_letter(column)].width = width
    result.conditional_formatting.add(
        "B5:B12",
        FormulaRule(
            formula=['OR(B5="Input needed",B5="Capacity Workshop Required")'],
            fill=PatternFill("solid", fgColor=LIGHT_RED),
        ),
    )
    result.conditional_formatting.add(
        "B5:B12",
        FormulaRule(
            formula=['OR(B5="Pilot / Light",B5="Small Production")'],
            fill=PatternFill("solid", fgColor=LIGHT_GREEN),
        ),
    )

    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.oddFooter.center.text = "NexoraAI — Customer Initial Discovery"
        ws.oddFooter.right.text = "Page &P of &N"
        ws.oddFooter.center.size = 8
        ws.oddFooter.center.color = GRAY

    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    wb.calculation.calcMode = "auto"
    return wb


def validate(path: Path) -> None:
    wb = load_workbook(path, data_only=False)
    assert wb.sheetnames == [
        "00 Start Here",
        "01 Basic Info",
        "02 HR Use Cases",
        "03 IT Ops Use Cases",
        "04 Data & Systems",
        "05 Security & Ops",
        "06 Simple Sizing",
        "07 Initial Summary",
    ]
    assert wb["07 Initial Summary"]["B6"].data_type == "f"
    assert wb["06 Simple Sizing"].max_row == 17
    assert sum(len(ws.tables) for ws in wb.worksheets) == 7


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    wb = build_workbook()
    wb.save(OUTPUT)
    validate(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
