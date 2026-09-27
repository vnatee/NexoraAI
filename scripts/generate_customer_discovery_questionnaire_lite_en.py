from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from generate_customer_discovery_questionnaire_lite import build_workbook


OUTPUT = Path(
    "docs/questionnaires/NexoraAI_Customer_Initial_Questionnaire_Lite_EN.xlsx"
)


TRANSLATIONS: dict[str, dict[str, str]] = {
    "00 Start Here": {
        "A2": "Initial discovery questionnaire for HR Employee Services and IT Operations | Estimated completion time: 30–45 minutes",
        "A4": "Purpose",
        "A5": "Use this workbook for the first customer discussion to understand business goals, target use cases, information sources, connected systems, user volume, and deployment preference before defining a demo, pilot, and initial budget.",
        "A7": "How to complete",
        "A8": "Complete the yellow cells. Enter TBC when an answer is not yet known. Select and prioritize the use cases of interest, then provide the sizing information currently available. Technical details can be confirmed during a later workshop.",
        "A10": "Security note",
        "A11": "Do not enter passwords, API keys, tokens, client secrets, connection strings, or real employee records. Record only the system name and responsible owner.",
        "A14": "The sizing in this workbook is an initial pilot and budgeting guide. Production sizing must be validated with the selected model, representative documents, integrations, and actual user workload.",
        "A17": "Customer",
    },
    "01 Basic Info": {
        "A2": "Initial questions to define the project goal, scope, and delivery approach",
        "B4": "Topic",
        "C4": "Question",
        "D4": "Example or hint",
        "E4": "Customer answer",
        "C5": "What problem do you want NexoraAI to solve first?",
        "D5": "For example: reduce HR response time or shorten incident investigation time.",
        "D6": "HR, IT Operations, Service Desk, or another selected business unit.",
        "C7": "Who will use the system?",
        "D7": "Employees, HR, Service Desk, NOC, system engineers, or managers.",
        "D8": "Target month or quarter.",
        "D9": "Azure, on-premises, hybrid, or not yet decided.",
        "C11": "How should users sign in?",
        "D12": "For example: HR can access only HR data, and IT can access only IT data.",
        "D13": "Thai, English, or both.",
        "D14": "For example: reduce search time by 50%, achieve 90% approved-answer accuracy, or reduce MTTR.",
    },
    "02 HR Use Cases": {
        "A2": "Select the use cases of interest for the initial demo or pilot. Start with Knowledge-based assistance and add HR system actions later if required.",
        "B4": "Use case",
        "C4": "Example question",
        "F4": "Current problem",
        "H4": "Expected result",
        "C5": "How many carried-over leave days can I use, and when do they expire?",
        "C6": "Am I eligible for this benefit, and what conditions apply?",
        "C7": "What information is required, and who must approve my request?",
        "C8": "How do I request an employment letter, and how long will it take?",
        "C9": "Can this course be reimbursed, and what approvals are required?",
        "C10": "What must a new employee complete during the first week?",
        "C11": "Route this request to the correct HR service team.",
        "C12": "Show my own information or submit my request after confirmation.",
        "C13": "Do not disclose another employee's salary or performance information.",
    },
    "03 IT Ops Use Cases": {
        "A2": "Select the use cases to address first and indicate whether the Agent should only read and recommend or perform an approved action.",
        "B4": "Use case",
        "C4": "Example question",
        "F4": "Current problem",
        "H4": "Expected result",
        "C5": "Which critical alerts are active, and which services are affected?",
        "C6": "Why has the trading application been slow during the last 30 minutes?",
        "C7": "Correlate APM, logs, database, infrastructure, and recent changes.",
        "C8": "Which approved runbook or known-error record matches this evidence?",
        "C9": "Create a timestamped timeline of events and supporting evidence.",
        "C10": "Create a chart of alerts by severity for the last three months.",
        "C11": "Prepare the information required to create an incident in ServiceNow or Jira.",
        "C12": "Recommend remediation with risks, validation steps, and rollback criteria.",
        "C13": "Draft a post-incident review from verified evidence.",
    },
    "04 Data & Systems": {
        "A2": "Initial information about documents, source systems, and connection methods",
        "B4": "Topic",
        "C4": "Question",
        "D4": "Example or hint",
        "E4": "Customer answer",
        "D7": "For example: approximately 5,000 files or 50 GB.",
        "D8": "PDF, Word, Excel, PowerPoint, scanned images, Thai, or English.",
        "D10": "For example: payroll is restricted to HR, and runbooks are restricted to IT.",
        "D11": "Record the system name and owner only; do not enter credentials.",
        "D15": "Record firewall, proxy, or private-network restrictions.",
        "D16": "Start with read-only access or require human approval for actions.",
    },
    "05 Security & Ops": {
        "A2": "Initial questions only. Detailed controls can be confirmed in a later Security and Architecture Workshop.",
        "B4": "Topic",
        "C4": "Question",
        "D4": "Example or hint",
        "E4": "Customer answer",
        "D5": "PII, salary, health, performance, customer information, or production logs.",
        "D6": "Identify prohibited information or role-restricted information.",
        "D7": "Create a ticket, change a system, send email, or submit an HR request.",
        "D8": "30, 90, 180, or 365 days, or another approved period.",
        "D11": "Provide RTO and RPO if known; otherwise enter TBC.",
    },
    "06 Simple Sizing": {
        "A2": "Complete the information currently available. These values provide an initial pilot and budget estimate only.",
        "B4": "Sizing question",
        "E5": "For example: 500",
        "E6": "For example: 100",
        "E7": "For example: 20",
        "E8": "For example: 5",
        "E9": "For example: 5,000",
        "E10": "For example: 50 GB",
        "E11": "For example: 200",
        "E12": "For example: 20%",
        "E13": "For example: 90 days",
        "E15": "Select TBC if not yet known.",
        "E17": "Consider growth in users, questions, and documents.",
    },
    "07 Initial Summary": {
        "A2": "Automatically generated initial summary after Simple Sizing is completed. Confirm the result through a Technical Workshop and representative load test.",
    },
}


def build_english_workbook():
    workbook = build_workbook()
    for sheet_name, cells in TRANSLATIONS.items():
        worksheet = workbook[sheet_name]
        for coordinate, value in cells.items():
            worksheet[coordinate] = value

    replacements = {
        "High / สูง": "High",
        "Medium / กลาง": "Medium",
        "Low / ต่ำ": "Low",
        "Later / ภายหลัง": "Later",
    }
    for worksheet in workbook.worksheets:
        for validation in worksheet.data_validations.dataValidation:
            if validation.formula1:
                for source, target in replacements.items():
                    validation.formula1 = validation.formula1.replace(source, target)
    return workbook


def validate(path: Path) -> None:
    workbook = load_workbook(path, data_only=False)
    thai = re.compile(r"[\u0E00-\u0E7F]")
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                assert not (
                    isinstance(cell.value, str) and thai.search(cell.value)
                ), f"Thai text remains in {worksheet.title}!{cell.coordinate}"
        for validation in worksheet.data_validations.dataValidation:
            assert not thai.search(validation.formula1 or "")
    assert workbook["07 Initial Summary"]["B7"].data_type == "f"
    assert sum(len(ws.tables) for ws in workbook.worksheets) == 7


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    workbook = build_english_workbook()
    workbook.save(OUTPUT)
    validate(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
