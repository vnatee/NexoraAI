#!/usr/bin/env python3
"""Build the stable NexoraAI Word manuals from their Markdown sources."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
LOGO = ROOT / "app" / "frontend" / "public" / "insight_sphere_new.png"

MANUALS = (
    (
        DOCS / "user-manual.md",
        DOCS / "NexoraAI-User-Manual.docx",
        "End users, business users, analysts, and team members",
        {
            "3. Workspace overview": ("Portal header and active Chat context", "Capture the header, account menu, Agent, model, Team, Data chips, and Change action."),
            "4. Chat sidebar": ("Chat workspace navigation", "Capture the expanded Workspace, Management, Insights, and Recent chats groups."),
            "5. Ask a question": ("Question composer", "Capture the composer, attachment action, Send action, and selected context without sensitive prompt content."),
            "6. Attachments": ("Attachment and Evidence matrix controls", "Capture safe sample attachments, processing status, Evidence matrix switch, and evidence-scope selector."),
            "7. Read and verify an answer": ("Answer, citations, table, and artifact controls", "Capture a sanitized response showing citations and a downloadable artifact or chart controls."),
            "9. Knowledge upload": ("Knowledge upload form", "Capture File, Data Group/scope, processing profile, and submission status using a non-sensitive sample file."),
            "10. My Usage": ("Personal usage dashboard", "Capture My Usage scope, date/bucket filters, and sanitized summary charts."),
            "11. Profile and password": ("My profile", "Capture the profile dialog with personal values replaced or blurred."),
        },
    ),
    (
        DOCS / "admin-manual.md",
        DOCS / "NexoraAI-Administrator-Manual.docx",
        "System Administrators, Team Administrators, security administrators, and platform operators",
        {
            "3. Open Administration": ("System Administration navigation", "Capture the collapsed menu groups and the SYSTEM ADMINISTRATION context header."),
            "4. Organization": ("Organization management", "Capture the Teams list and Users & permissions navigation using sanitized identities."),
            "5. Team workspace": ("Team workspace navigation", "Capture a Team context header and the Members, Roles & permissions, Data access, Agents, and Team integrations menus."),
            "6. Agents": ("Agent configuration", "Capture Agent Info and Integrations tabs with a safe example Agent and no prompt containing customer data."),
            "7. Unified integrations": ("Configured connections and Browse catalog", "Capture both tabs and connection-type filters; hide URLs, tenant IDs, and identifiers."),
            "8. Governance": ("Governance controls", "Capture Guardrail policy, Usage quotas, or Chat history retention with safe example values."),
            "9. LLM providers": ("LLM provider configuration", "Capture the provider list and Add/Edit dialog; hide endpoint tenant names and all credentials."),
            "10. System health": ("On-demand System Health", "Capture runtime resources and Core services, including HTTPS gateway · Nginx when TLS is enabled."),
            "11. License": ("License status", "Capture status, Installation ID with most characters blurred, Trial/license dates, and entitlement chips. Never capture the raw key."),
            "12. Sign-in & authentication": ("Authentication settings", "Capture enabled sign-in methods with IDs blurred and the client-secret field empty."),
            "13. Knowledge ingestion": ("Ingestion and processing status", "Capture a sanitized job list showing normal lifecycle states."),
        },
    ),
)

BLUE = "1F4E78"
LIGHT_BLUE = "D9EAF7"
PALE_BLUE = "EEF5FA"
GRAY = "666666"
LIGHT_GRAY = "F2F2F2"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "A6A6A6", size: str = "8") -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = properties.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        properties.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def add_field(paragraph, instruction: str) -> None:
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph._p.extend((begin, text, separate, end))


def add_hyperlink(paragraph, label: str, target: str) -> None:
    relationship = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    properties.extend((color, underline))
    text = OxmlElement("w:t")
    text.text = label
    run.extend((properties, text))
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")


def add_inline(paragraph, value: str) -> None:
    cursor = 0
    for match in INLINE.finditer(value):
        if match.start() > cursor:
            paragraph.add_run(value[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            paragraph.add_run(token[2:-2]).bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(80, 80, 80)
        else:
            label, target = re.match(r"\[([^]]+)\]\(([^)]+)\)", token).groups()
            add_hyperlink(paragraph, label, target)
        cursor = match.end()
    if cursor < len(value):
        paragraph.add_run(value[cursor:])


def configure_document(document: Document, title: str) -> None:
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    document.settings._element.append(update_fields)
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.75)
    section.right_margin = Inches(0.75)

    normal = document.styles["Normal"]
    normal.font.name = "Liberation Sans"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE

    for name, size, color in (
        ("Title", 30, BLUE),
        ("Heading 1", 19, BLUE),
        ("Heading 2", 15, BLUE),
        ("Heading 3", 12, GRAY),
    ):
        style = document.styles[name]
        style.font.name = "Liberation Sans"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)

    header = section.header.paragraphs[0]
    header.text = f"NexoraAI  |  {title}"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("NexoraAI  •  Internal Product Documentation  •  Page ")
    add_field(footer, " PAGE ")
    footer.add_run(" of ")
    add_field(footer, " NUMPAGES ")
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(GRAY)


def add_cover(document: Document, title: str, audience: str) -> None:
    if LOGO.is_file():
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(LOGO), width=Inches(1.35))
    paragraph = document.add_paragraph(style="Title")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run(title)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Product Manual")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    document.add_paragraph()
    for label, value in (
        ("Audience", audience),
        ("Language", "English"),
        ("Last verified", "8 September 2026"),
        ("Document status", "Controlled working manual — update in place"),
    ):
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run(f"{label}: ").bold = True
        paragraph.add_run(value)
    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(24)
    run = note.add_run("Screenshot placeholders are intentionally blank. Insert images from the approved delivery environment and remove or blur sensitive values.")
    run.italic = True
    run.font.color.rgb = RGBColor.from_string(GRAY)
    document.add_page_break()

    heading = document.add_paragraph("Table of Contents", style="Heading 1")
    heading.paragraph_format.space_after = Pt(12)
    toc = document.add_paragraph()
    add_field(toc, ' TOC \\o "1-3" \\h \\z \\u ')
    hint = document.add_paragraph("In Microsoft Word: right-click the table and select Update Field → Update entire table.")
    hint.runs[0].italic = True
    hint.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    document.add_page_break()


def add_screenshot_placeholder(document: Document, number: int, title: str, instruction: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.65)
    row = table.rows[0]
    row.height = Inches(2.35)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    cell = row.cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_shading(cell, PALE_BLUE)
    set_cell_border(cell, BLUE, "12")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(f"SCREENSHOT PLACEHOLDER {number}\n\n{title}\n\n")
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    detail = paragraph.add_run(instruction)
    detail.italic = True
    detail.font.size = Pt(9)
    detail.font.color.rgb = RGBColor.from_string(GRAY)
    caption = document.add_paragraph(f"Figure {number} — {title}")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.runs[0].italic = True
    caption.runs[0].font.size = Pt(9)
    caption.runs[0].font.color.rgb = RGBColor.from_string(GRAY)


def table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def is_separator(line: str) -> bool:
    return bool(re.fullmatch(r"\|?[ :|-]+\|?", line.strip())) and "-" in line


def add_markdown_table(document: Document, lines: list[str]) -> None:
    rows = [table_row(line) for line in lines if not is_separator(line)]
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.autofit = True
    for row_index, values in enumerate(rows):
        for column_index in range(width):
            cell = table.cell(row_index, column_index)
            set_cell_border(cell)
            if row_index == 0:
                set_cell_shading(cell, BLUE)
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, values[column_index] if column_index < len(values) else "")
            for run in paragraph.runs:
                run.font.size = Pt(8.5)
                if row_index == 0:
                    run.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code_block(document: Document, lines: list[str]) -> None:
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, LIGHT_GRAY)
    set_cell_border(cell, "D9D9D9", "4")
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"
    run.font.size = Pt(8.5)


def render_markdown(document: Document, source: Path, placeholders: dict[str, tuple[str, str]]) -> None:
    lines = source.read_text(encoding="utf-8").splitlines()
    figure = 0
    index = 0
    in_code = False
    code_lines: list[str] = []
    skipped_metadata = False
    while index < len(lines):
        line = lines[index].rstrip()
        stripped = line.strip()
        if stripped.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
            in_code = not in_code
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue
        if index == 0 and stripped.startswith("# "):
            index += 1
            continue
        if not skipped_metadata and (not stripped or stripped.startswith("**Language:") or stripped.startswith("**Audience:") or stripped.startswith("**Last verified:") or stripped.startswith("**Applies to:")):
            index += 1
            continue
        skipped_metadata = True
        if not stripped:
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines) and is_separator(lines[index + 1]):
            table_lines = [line, lines[index + 1]]
            index += 2
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index])
                index += 1
            add_markdown_table(document, table_lines)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            level = len(heading.group(1)) - 1
            text = heading.group(2)
            if level == 1 and any(paragraph.text for paragraph in document.paragraphs[-3:]):
                document.add_page_break()
            document.add_heading(text, level=level)
            if text in placeholders:
                figure += 1
                add_screenshot_placeholder(document, figure, *placeholders[text])
            index += 1
            continue
        unordered = re.match(r"^\s*-\s+(.+)$", line)
        ordered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if unordered or ordered:
            paragraph = document.add_paragraph(style="List Bullet" if unordered else "List Number")
            add_inline(paragraph, (unordered or ordered).group(1))
            index += 1
            continue
        if stripped.startswith(">"):
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Inches(0.3)
            paragraph.paragraph_format.right_indent = Inches(0.2)
            paragraph.paragraph_format.space_before = Pt(4)
            paragraph.paragraph_format.space_after = Pt(8)
            add_inline(paragraph, stripped.lstrip("> "))
            for run in paragraph.runs:
                run.italic = True
                run.font.color.rgb = RGBColor.from_string(GRAY)
            index += 1
            continue
        paragraph_lines = [stripped]
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if not candidate or candidate.startswith(("#", "- ", ">", "```", "|")) or re.match(r"^\d+\.\s", candidate):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        add_inline(paragraph, " ".join(paragraph_lines))


def build(source: Path, output: Path, audience: str, placeholders: dict[str, tuple[str, str]]) -> None:
    source_text = source.read_text(encoding="utf-8")
    missing_headings = [heading for heading in placeholders if f"## {heading}" not in source_text]
    if missing_headings:
        raise ValueError(f"Screenshot placeholder headings are missing from {source.name}: {missing_headings}")
    title = source_text.splitlines()[0].removeprefix("# ").strip()
    document = Document()
    configure_document(document, title)
    add_cover(document, title, audience)
    render_markdown(document, source, placeholders)
    properties = document.core_properties
    properties.title = title
    properties.subject = "NexoraAI product usage and administration"
    properties.author = "NexoraAI"
    properties.keywords = "NexoraAI, manual, user guide, administrator guide"
    document.save(output)


def main() -> int:
    for source, output, audience, placeholders in MANUALS:
        build(source, output, audience, placeholders)
        print(f"Created {output.relative_to(ROOT)}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
