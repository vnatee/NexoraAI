from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from generate_product_sales_deck import (
    BLUE,
    CYAN,
    GREEN,
    INK,
    LIGHT,
    MUTED,
    NAVY,
    NAVY_2,
    ORANGE,
    PALE_BLUE,
    PALE_CYAN,
    PALE_ORANGE,
    RED,
    WHITE,
    OUTPUT as TH_OUTPUT,
    add_bullets,
    add_card,
    add_footer,
    add_header,
    add_rect,
    add_text,
    build_deck,
    new_slide,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "presentations" / "InsightSphere_Product_Sales_Deck_EN.pptx"


TRANSLATIONS = {
    "เปลี่ยนความรู้และระบบงานขององค์กร\nให้เป็น AI ที่ค้นหา วิเคราะห์ และลงมือทำได้":
        "Turn enterprise knowledge and workflows\ninto AI that can search, reason, and act",
    "AI Platform เดียว สำหรับความรู้ คน และระบบงาน":
        "One AI Platform for Knowledge, People, and Workflows",
    "InsightSphere ช่วยให้องค์กรสร้าง AI Agent ที่ตอบจากข้อมูลจริง เชื่อมต่อระบบเดิม\nและควบคุมสิทธิ์ได้ โดยเลือกใช้ Cloud, Local หรือ Hybrid ได้ตามนโยบาย":
        "Build AI agents grounded in trusted enterprise data, connected to existing systems,\nwith governed access and the freedom to deploy in cloud, local, or hybrid environments.",
    "ค้นหาและอ้างอิงความรู้": "Search and Cite Knowledge",
    "RAG พร้อม Citation ระดับไฟล์/หน้า และขอบเขตข้อมูล Personal, Team, Global":
        "RAG with file- and page-level citations across Personal, Team, and Global data scopes.",
    "วิเคราะห์หลายรูปแบบ": "Analyze Multiple Content Types",
    "ข้อความ เอกสาร ภาพ OCR และ Evidence Matrix สำหรับงานตรวจข้อกำหนด":
        "Text, documents, images, OCR, and evidence matrices for requirements analysis.",
    "Agent เชื่อมต่อและลงมือทำ": "Connect Agents to Action",
    "เรียก Tool, MCP, Database, API, Mail และ Messaging ภายใต้นโยบายองค์กร":
        "Invoke tools, MCP, databases, APIs, mail, and messaging under enterprise policy.",
    "ผลลัพธ์: จาก “ทดลองใช้ AI” สู่ “ระบบงาน AI ที่บริหารจัดการได้”":
        "Outcome: Move from AI experiments to governed, operational AI workflows.",
    "ความท้าทายที่องค์กรพบเมื่อเริ่มใช้ Generative AI":
        "The Enterprise Challenges of Adopting Generative AI",
    "ข้อมูลกระจัดกระจาย": "Fragmented Knowledge",
    "เอกสารอยู่หลายทีม หลายระบบ และค้นหาคำตอบที่เชื่อถือได้ยาก":
        "Documents span teams and systems, making trusted answers difficult to find.",
    "คำตอบไม่มีหลักฐาน": "Answers Without Evidence",
    "AI ตอบได้ แต่ผู้ใช้ตรวจสอบแหล่งข้อมูลและหน้าที่อ้างอิงไม่ได้":
        "AI can respond, but users cannot verify the source file or cited page.",
    "สิทธิ์เข้าถึงซับซ้อน": "Complex Access Control",
    "ข้อมูลแต่ละทีมมีระดับความลับต่างกัน และต้องป้องกันการเห็นข้อมูลข้ามทีม":
        "Teams have different confidentiality needs and cross-team exposure must be prevented.",
    "เชื่อมต่อระบบยาก": "Disconnected from Workflows",
    "AI ที่ตอบอย่างเดียวไม่ลดขั้นตอน หากยังเรียก API, Database หรือ Tool ไม่ได้":
        "Answers alone do not streamline work if AI cannot call APIs, databases, or tools.",
    "ข้อกำหนดด้านต้นทุน ข้อมูล และโมเดลเปลี่ยนเร็ว แต่ระบบผูกกับผู้ให้บริการรายเดียว":
        "Cost, data, and model requirements change quickly while the platform remains locked to one vendor.",
    "ขาด Audit, Policy, Usage และวิธีบริหาร Agent/Model แบบศูนย์กลาง":
        "Organizations lack centralized audit, policy, usage, agent, and model management.",
    "แพลตฟอร์มที่รวม Knowledge, Agent และ Governance":
        "A Unified Platform for Knowledge, Agents, and Governance",
    "Enterprise Knowledge ที่ตอบได้และตรวจสอบได้":
        "Enterprise Knowledge that Answers — and Proves It",
    "รับข้อมูลหลายขอบเขต": "Multiple Data Scopes",
    "Personal สำหรับผู้ใช้, Team + Data Group สำหรับงานร่วมกัน และ Global สำหรับความรู้ส่วนกลาง":
        "Personal knowledge for individuals, Team and Data Group access for collaboration, and Global knowledge for everyone.",
    "Routing ตามชนิดไฟล์ ทีม และลำดับความสำคัญ พร้อม Local parser, Docling และผู้ให้บริการ OCR":
        "Route by file type, team, and priority using local parsers, Docling, or managed OCR providers.",
    "Citation และ Page Evidence": "Citations and Page Evidence",
    "คำตอบเชื่อมกลับไปยังไฟล์และหน้าต้นทาง พร้อม authorization ก่อนเปิดเอกสารทุกครั้ง":
        "Answers link to source files and pages, with authorization enforced every time a document is opened.",
    "วงจรข้อมูลที่ควบคุมได้": "A Governed Knowledge Lifecycle",
    "ลดเวลาค้นหา • เพิ่มความน่าเชื่อถือ • รักษาสิทธิ์ข้อมูล • รองรับการตรวจสอบย้อนหลัง":
        "Reduce search time • Improve trust • Preserve access control • Support auditability",
    "AI Agent ที่ไม่ได้แค่ตอบ แต่ทำงานร่วมกับระบบองค์กร":
        "AI Agents that Do More than Answer",
    "กำหนด Agent แยกตามบทบาท พร้อม Model, Prompt, Data Scope และ Tool ที่เหมาะกับงาน":
        "Configure role-specific agents with the right model, prompt, data scope, and tools.",
    "ตอบนโยบายพนักงาน\nจากเอกสารที่ได้รับสิทธิ์":
        "Answer employee policy questions\nfrom authorized documents",
    "ค้นข้อมูลผลิตภัณฑ์\nและสร้างคำตอบเชิงหลักฐาน":
        "Find product knowledge\nand create evidence-based responses",
    "ตรวจ Monitor / Log\nผ่าน MCP และสรุปเหตุการณ์":
        "Inspect monitors and logs\nthrough MCP and summarize incidents",
    "Query ฐานข้อมูลแบบ read-only\nและสร้าง Chart":
        "Query databases in read-only mode\nand generate charts",
    "Tool policy จำกัดจำนวนครั้ง เวลา และแหล่งที่เรียกใช้":
        "Tool policies limit call volume, execution time, and permitted sources.",
    "Sub-agent delegation พร้อมป้องกัน cycle และ depth เกินกำหนด":
        "Sub-agent delegation includes cycle prevention and bounded depth.",
    "Audit event เชื่อม model, tool และ trace โดยไม่เก็บ raw secrets":
        "Audit events connect models, tools, and traces without storing raw secrets.",
    "จากเอกสารข้อกำหนด สู่ Evidence Matrix ที่ตรวจสอบได้":
        "From Requirements Documents to an Auditable Evidence Matrix",
    "ข้อกำหนดงาน\nTOR / RFP / Policy\nChecklist / Standard":
        "Requirements\nTOR / RFP / Policy\nChecklist / Standard",
    "รองรับเอกสารแนบชั่วคราว\nหรือ Knowledge ที่องค์กรจัดเก็บ":
        "Use temporary attachments\nor governed enterprise Knowledge",
    "ระบบต้องมี Audit": "The system must provide audit trails",
    "พบ": "Found",
    "รองรับ Local": "Support local deployment",
    "ข้อกำหนดเฉพาะ": "Customer-specific requirement",
    "ไม่พบ": "Not found",
    "Benefit: ตรวจข้อกำหนดได้เป็นระบบ เห็นทั้งหลักฐานและช่องว่าง ลดความเสี่ยงจากคำตอบที่ไม่มีข้อมูลรองรับ":
        "Benefit: Evaluate requirements systematically, expose evidence and gaps, and reduce unsupported conclusions.",
    "Security & Governance เป็นส่วนหนึ่งของ Platform":
        "Security and Governance Built into the Platform",
    "Local User + Microsoft Entra ID\nรองรับ bootstrap super admin":
        "Local users + Microsoft Entra ID\nwith secure super-admin bootstrap",
    "Team, Data Group, Personal/Global ACL\nตรวจสิทธิ์ทั้งค้นหาและเปิด Citation":
        "Team, Data Group, Personal/Global ACLs\nenforced for retrieval and citation access",
    "Fernet encryption สำหรับ provider key\nไม่ส่ง secret กลับ frontend":
        "Fernet encryption for provider keys\nsecrets are never returned to the frontend",
    "Trace model/tool activity\nพร้อม metadata โดยไม่เก็บ raw prompt/output":
        "Trace model and tool activity\nwithout persisting raw prompts or outputs",
    "History และ attachment แยกตาม user\nconversation และ data scope":
        "History and attachments isolated by user,\nconversation, and data scope",
    "Designed for enterprise control — ไม่ลดทอน Security Contract เพื่อความสะดวก":
        "Designed for enterprise control — security contracts are never weakened for convenience.",
    "เลือก Model และ Provider ให้เหมาะกับแต่ละ Agent":
        "Choose the Right Model and Provider for Every Agent",
    "ไม่บังคับให้ทุกงานใช้โมเดลเดียว — แยก Chat Model, Embedding และ Processing Provider ได้":
        "Avoid a one-model-fits-all approach — separate chat, embedding, and processing providers.",
    "Portal ทดสอบ Tool Calling และ Vision capability ก่อนเปิดใช้งาน":
        "The portal tests tool-calling and vision capabilities before activation.",
    "กำหนด Model ต่อ Agent เพื่อสมดุล Cost, Latency, Privacy และ Quality":
        "Select models per agent to balance cost, latency, privacy, and quality.",
    "เปลี่ยน provider ได้โดยรักษา Agent, ACL และ Knowledge architecture เดิม":
        "Change providers while preserving agents, ACLs, and the Knowledge architecture.",
    "Deployment ที่ยืดหยุ่น: Azure, Local หรือ Hybrid":
        "Flexible Deployment: Azure, Local, or Hybrid",
    "Local data plane\nเลือก Cloud LLM\nหรือ Local Model\nตาม Agent / Policy":
        "Local data plane\nCloud or local LLMs\nselected per agent\nand policy",
    "แยก Provider Interface ชัดเจน — รักษา Azure behavior พร้อมเปิดทางสู่ Self-hosted":
        "Clean provider interfaces preserve Azure behavior while enabling self-hosted deployment.",
    "ตัวอย่าง Use Case ที่สร้างมูลค่าได้เร็ว": "High-Value Use Cases to Start With",
    "ตอบคำถามนโยบาย คู่มือ และความรู้ภายใน พร้อม Citation":
        "Answer policy, manual, and internal knowledge questions with citations.",
    "เปรียบเทียบข้อกำหนดกับหลักฐานใน Knowledge และแสดง Gap":
        "Compare requirements with Knowledge evidence and identify gaps.",
    "ค้น Monitor/Log ผ่าน MCP สรุปเหตุการณ์ และแนะนำขั้นตอนแก้ไข":
        "Search monitors and logs through MCP, summarize incidents, and recommend next steps.",
    "Query ฐานข้อมูลแบบควบคุมสิทธิ์ สรุปและสร้างกราฟ":
        "Query governed databases, summarize results, and generate charts.",
    "ค้น Product Knowledge สร้างคำตอบที่ตรวจสอบแหล่งข้อมูลได้":
        "Find product knowledge and create source-verifiable responses.",
    "เรียก API, Mail, Messaging และ Sub-agent ตาม Tool policy":
        "Invoke APIs, mail, messaging, and sub-agents under tool policy.",
    "Business Benefits ที่เชื่อมโยงกับการวัดผล":
        "Business Benefits Connected to Measurable Outcomes",
    "ลดเวลาค้นหา สรุป และจัดทำคำตอบซ้ำ":
        "Reduce time spent searching, summarizing, and recreating answers.",
    "วัด: เวลาเฉลี่ยต่อเคส": "Measure: average time per case",
    "คำตอบมีหลักฐานและเห็นช่องว่างข้อมูล":
        "Deliver evidence-backed answers and make knowledge gaps visible.",
    "วัด: acceptance / rework rate": "Measure: acceptance and rework rate",
    "สิทธิ์ Agent Tool และข้อมูลบริหารจากศูนย์กลาง":
        "Manage data, agent, and tool permissions centrally.",
    "วัด: policy violation / audit coverage": "Measure: policy violations and audit coverage",
    "เปลี่ยน Model และ Deployment ตามงานได้":
        "Choose models and deployment modes according to workload needs.",
    "วัด: cost per task / latency": "Measure: cost per task and latency",
    "เริ่มจาก Baseline ของลูกค้า แล้ววัดผล Pilot ก่อนสรุป ROI — ไม่ใช้ตัวเลขสมมติ":
        "Establish the customer baseline and validate pilot results before claiming ROI — no invented numbers.",
    "แนวทางเริ่มใช้งานแบบลดความเสี่ยง": "A Low-Risk Path to Adoption",
    "เลือก 1–2 Use Case\nกำหนด Data/ACL/KPI":
        "Select 1–2 use cases\nDefine data, ACLs, and KPIs",
    "เชื่อมข้อมูลจริง\nตั้ง Agent และ Guardrails":
        "Connect real data\nConfigure agents and guardrails",
    "Pilot ใช้ข้อมูลจริงในขอบเขตจำกัด พร้อมผู้รับผิดชอบข้อมูล":
        "Run the pilot with scoped production-like data and named data owners.",
    "Production scope รวม TLS, Backup/Restore, Monitoring, Capacity และ Support model":
        "Production scope includes TLS, backup/restore, monitoring, capacity, and a support model.",
    "ทำไม InsightSphere เหมาะกับ Enterprise AI":
        "Why InsightSphere for Enterprise AI",
    "Citation, Page, Evidence และ Gap ถูกออกแบบเป็น workflow ไม่ใช่ส่วนเสริม":
        "Citations, page evidence, and gaps are designed into the workflow — not added later.",
    "Team/Data Group ACL อยู่ใน ingestion, retrieval และ citation authorization":
        "Team and Data Group ACLs span ingestion, retrieval, and citation authorization.",
    "Agent เรียก Tool ได้จริงภายใต้ policy, permission และ audit":
        "Agents invoke real tools under policy, permission, and audit controls.",
    "Cloud, Local และ Hybrid เปลี่ยนตามข้อกำหนดโดยไม่เขียน platform ใหม่":
        "Move between cloud, local, and hybrid providers without rebuilding the platform.",
    "Admin Portal, model capability test, persistent jobs และ recovery":
        "Admin portal, model capability testing, persistent jobs, and recovery.",
    "MCP, LangGraph, sub-agent และ provider adapter รองรับการเติบโต":
        "MCP, LangGraph, sub-agents, and provider adapters support future growth.",
    "เริ่มจาก Use Case ที่เห็นผล\nแล้วขยายเป็น Enterprise AI Platform":
        "Start with a High-Value Use Case\nThen Scale into an Enterprise AI Platform",
    "ข้อเสนอสำหรับ Next Step": "Recommended Next Steps",
    "Workshop: Use Case, Data Scope, Security และ KPI":
        "Workshop: Use cases, data scope, security, and KPIs",
    "Technical Assessment: Deployment และ Integration":
        "Technical assessment: deployment and integrations",
    "Pilot Plan: ขอบเขต ผู้ใช้ ระยะเวลา และเกณฑ์ความสำเร็จ":
        "Pilot plan: scope, users, timeline, and success criteria",
    "การประเมิน Capacity ต้องเริ่มจาก Concurrent Workload":
        "Capacity Planning Starts with Concurrent Workload",
    "จำนวน Account ไม่เท่ากับจำนวนงาน AI ที่ประมวลผลพร้อมกัน":
        "Account count is not the same as simultaneous AI workload.",
    "ผู้ใช้ทั้งหมด": "Total users",
    "ใช้วางแผน Identity\nและข้อมูลระยะยาว":
        "Plan identity and\nlong-term data volume",
    "ผู้ใช้ที่ใช้งานจริง": "Active users",
    "ดู Daily/Monthly Active\nและรูปแบบการใช้งาน":
        "Measure daily/monthly activity\nand usage patterns",
    "คำขอที่กำลังทำงาน": "In-flight requests",
    "กำหนด App replicas,\nconnection และ queue":
        "Size app replicas,\nconnections, and queues",
    "LLM ที่สร้างคำตอบพร้อมกัน": "Concurrent LLM generations",
    "กำหนด GPU slots\nหรือ Cloud quota":
        "Size GPU slots\nor cloud quota",
    "Capacity โดยประมาณ = Peak concurrent requests × เวลาประมวลผล × ขนาด Context / Tool workload":
        "Estimated capacity = peak concurrent requests × processing time × context and tool workload",
    "ตัวเลขในเอกสารนี้เป็น Planning baseline — ต้องยืนยันด้วย Load Test จาก use case และ model จริง":
        "These figures are planning baselines—validate them with the actual use case, model, and load test.",
    "Hardware Sizing Baseline สำหรับ Local / Hybrid":
        "Hardware Sizing Baseline for Local and Hybrid Deployments",
    "ระดับ": "Stage",
    "แนวทางรองรับเบื้องต้น": "Indicative capacity",
    "SSD 200 GB\nไม่บังคับ GPU": "SSD 200 GB\nGPU optional",
    "1–5 ผู้ใช้พร้อมกัน\nเหมาะกับทดสอบและสาธิต":
        "1–5 concurrent users\nDevelopment and demos",
    "NVMe 500 GB\nLLM ใช้ Cloud": "NVMe 500 GB\nCloud-hosted LLM",
    "ประมาณ 10–30 concurrent Chat\nขึ้นกับ provider quota และ tools":
        "Approx. 10–30 concurrent chats\nDepends on provider quota and tools",
    "1–2 local generations พร้อมกัน\nเหมาะกับ light / staggered usage":
        "1–2 concurrent local generations\nLight or staggered usage",
    "App 2+ nodes\nnode ละ 4–8 vCPU / 16 GB":
        "2+ app nodes\n4–8 vCPU / 16 GB each",
    "แยก Stateful services\nและ GPU inference node":
        "Dedicated stateful services\nand GPU inference nodes",
    "ประมาณ 30–100 concurrent users\nหลังผ่าน Load Test":
        "Approx. 30–100 concurrent users\nafter load-test validation",
    "หมายเหตุ: พื้นที่จัดเก็บต้องคำนวณเพิ่มจากเอกสารต้นฉบับ + vector index + chat/attachment retention + backup":
        "Storage must include source documents, vector indexes, chat and attachment retention, and backups.",
    "จาก Single-node Pilot สู่ Production Scale":
        "From a Single-Node Pilot to Production Scale",
    "Compose เครื่องเดียว — ไม่มี automatic failover":
        "Single-host Compose with no automatic failover",
    "Hypercorn application process เดียว":
        "Single Hypercorn application process",
    "Ollama parallel = 1 เป็นคอขวด Local LLM":
        "Ollama parallel = 1 limits local LLM throughput",
    "MongoDB, MinIO และ Qdrant เป็น single instance":
        "MongoDB, MinIO, and Qdrant are single instances",
    "Ingestion อยู่ใน App และจำกัด 3 งานต่อ process":
        "Ingestion runs in the app, limited to three jobs per process",
    "Load balancer และ App 2+ replicas":
        "Load balancer with 2+ app replicas",
    "Queue แยก OCR / ingestion / cleanup":
        "Dedicated queues for OCR, ingestion, and cleanup",
    "Cloud quota หรือ GPU pool / vLLM":
        "Cloud quota or a GPU pool with vLLM",
    "Scale trigger: p95 latency, queue depth, provider throttling, GPU utilization และ ingestion backlog":
        "Scale triggers: p95 latency, queue depth, provider throttling, GPU utilization, and ingestion backlog",
}


def normalize(value: str) -> str:
    return value.strip().replace("\v", "\n")


def add_feature_slides(prs: Presentation) -> None:
    """Add English-first feature slides before the original closing slide."""
    feature_slide_ids = []

    # 15 — Guardrails and usage governance
    slide = new_slide(prs, LIGHT)
    feature_slide_ids.append(prs.slides._sldIdLst[-1])
    add_header(slide, "Enterprise AI Guardrails and Usage Governance", "Feature Update", 15)
    add_text(slide, "Policy inheritance keeps every team and agent within the enterprise minimum.",
             Inches(0.72), Inches(1.86), Inches(11.85), Inches(0.45), 18, INK, True)
    hierarchy = [
        ("SYSTEM", "Enterprise minimum", "Cannot be weakened downstream", BLUE),
        ("TEAM", "Team-specific controls", "May enforce stricter policy", CYAN),
        ("AGENT / USER", "Workload-level limits", "Narrower settings and quotas", ORANGE),
    ]
    for i, (badge, title, body, color) in enumerate(hierarchy):
        x = 0.72 + i * 4.08
        add_card(slide, title, body, Inches(x), Inches(2.55), Inches(3.65), Inches(1.42), color, WHITE, badge)
        if i < 2:
            add_text(slide, "→", Inches(x + 3.65), Inches(2.94), Inches(0.42), Inches(0.4), 20, MUTED, True, PP_ALIGN.CENTER)
    controls = [
        ("GUARDRAILS", "Prompt injection • Secrets • PII/DLP • Content safety", RED),
        ("USAGE QUOTAS", "Tokens • Provider/model rules • Document pages • Requests/min", GREEN),
        ("PRIVACY-SAFE AUDIT", "Trace and enforcement metadata without raw prompts or outputs", BLUE),
    ]
    for i, (title, body, color) in enumerate(controls):
        x = 0.72 + i * 4.08
        add_rect(slide, Inches(x), Inches(4.5), Inches(3.65), Inches(1.35), WHITE, True, color)
        add_text(slide, title, Inches(x + 0.22), Inches(4.72), Inches(3.21), Inches(0.32), 11, color, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x + 0.26), Inches(5.18), Inches(3.13), Inches(0.48), 11, INK, False, PP_ALIGN.CENTER)
    add_text(slide, "Atomic reservations and idempotent metering reduce quota overshoot across application replicas.",
             Inches(0.76), Inches(6.36), Inches(11.8), Inches(0.38), 13, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 16 — Team-scoped integrations and MCP
    slide = new_slide(prs, NAVY)
    feature_slide_ids.append(prs.slides._sldIdLst[-1])
    add_header(slide, "Team-Scoped Integration and MCP Catalog", "Feature Update", 16, True)
    add_text(slide, "Multiple connections per provider—isolated by Team and assignable only to authorized Agents.",
             Inches(0.72), Inches(1.85), Inches(11.85), Inches(0.45), 18, WHITE, True)
    flow = [
        ("ADMIN CATALOG", "Create as Unassigned", BLUE),
        ("TEAM OWNERSHIP", "Assign once to one Team", CYAN),
        ("AGENT BINDING", "Select approved tools", ORANGE),
        ("RUNTIME", "Reject cross-Team access", GREEN),
    ]
    for i, (title, body, color) in enumerate(flow):
        x = 0.72 + i * 3.08
        add_rect(slide, Inches(x), Inches(2.62), Inches(2.68), Inches(1.18), RGBColor(24, 47, 72), True, color)
        add_text(slide, title, Inches(x + 0.18), Inches(2.83), Inches(2.32), Inches(0.3), 10, color, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x + 0.18), Inches(3.27), Inches(2.32), Inches(0.3), 11, WHITE, False, PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, "→", Inches(x + 2.68), Inches(2.96), Inches(0.38), Inches(0.38), 18, CYAN, True, PP_ALIGN.CENTER)
    integrations = [
        ("DATADOG", "Multiple tenants and sites\nBearer or API/Application keys", ORANGE),
        ("MICROSOFT AZURE", "Managed local or external MCP\nRead-only inventory tools", BLUE),
        ("DRAW.IO", "Hosted MCP connection\nDownloadable diagram artifacts", CYAN),
        ("ENTERPRISE SYSTEMS", "REST • Database • Mail\nMessaging • Sub-agents", GREEN),
    ]
    for i, (title, body, color) in enumerate(integrations):
        x = 0.72 + i * 3.08
        add_rect(slide, Inches(x), Inches(4.35), Inches(2.68), Inches(1.55), RGBColor(24, 47, 72), True, RGBColor(48, 78, 107))
        add_rect(slide, Inches(x + 0.2), Inches(4.58), Inches(0.1), Inches(1.08), color, False)
        add_text(slide, title, Inches(x + 0.43), Inches(4.55), Inches(2.02), Inches(0.3), 10, color, True)
        add_text(slide, body, Inches(x + 0.43), Inches(5.02), Inches(2.02), Inches(0.65), 10, WHITE)
    add_text(slide, "Encrypted credentials • Tool discovery • Allowlists • Timeouts • Server-enforced ownership",
             Inches(0.75), Inches(6.4), Inches(11.85), Inches(0.34), 12, CYAN, True, PP_ALIGN.CENTER)
    add_footer(slide, dark=True)

    # 17 — Azure inventory to Draw.io
    slide = new_slide(prs)
    feature_slide_ids.append(prs.slides._sldIdLst[-1])
    add_header(slide, "Azure Inventory to Draw.io Architecture", "Feature Update", 17)
    add_text(slide, "Turn live Azure inventory into an editable architecture artifact—with evidence preserved.",
             Inches(0.72), Inches(1.84), Inches(11.85), Inches(0.5), 18, INK, True)
    stages = [
        ("AZURE TENANT", "Subscription and\nresource inventory", BLUE),
        ("AZURE MCP", "Managed local container\nRead-only tool policy", CYAN),
        ("AZURE ENGINEER", "Classify resources\nand relationships", ORANGE),
        ("DRAW.IO MCP", "Generate valid XML\nwith Azure icons", GREEN),
        ("ARTIFACT", "Download and edit\n.drawio file", BLUE),
    ]
    for i, (title, body, color) in enumerate(stages):
        x = 0.58 + i * 2.55
        add_rect(slide, Inches(x), Inches(2.75), Inches(2.1), Inches(2.0), LIGHT, True, color)
        add_rect(slide, Inches(x + 0.67), Inches(2.98), Inches(0.76), Inches(0.48), color, True)
        add_text(slide, str(i + 1), Inches(x + 0.67), Inches(2.98), Inches(0.76), Inches(0.48), 13, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, title, Inches(x + 0.15), Inches(3.67), Inches(1.8), Inches(0.35), 12, color, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x + 0.18), Inches(4.12), Inches(1.74), Inches(0.5), 10, INK, False, PP_ALIGN.CENTER)
        if i < 4:
            add_text(slide, "→", Inches(x + 2.1), Inches(3.54), Inches(0.43), Inches(0.45), 19, MUTED, True, PP_ALIGN.CENTER)
    add_rect(slide, Inches(0.72), Inches(5.35), Inches(11.85), Inches(0.92), PALE_BLUE, True, BLUE)
    add_text(slide, "Solid connectors = evidenced relationships  •  Dashed connectors = explicitly labelled inferred relationships",
             Inches(0.95), Inches(5.58), Inches(11.4), Inches(0.34), 14, BLUE, True, PP_ALIGN.CENTER)
    add_text(slide, "Current scope is intentionally read-only: inventory and diagram generation do not provision or modify Azure resources.",
             Inches(0.78), Inches(6.47), Inches(11.75), Inches(0.3), 11, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide)

    # 18 — Durable operations
    slide = new_slide(prs, LIGHT)
    feature_slide_ids.append(prs.slides._sldIdLst[-1])
    add_header(slide, "Durable Operations: History, Ingestion, and Recovery", "Feature Update", 18)
    feature_cards = [
        ("DATABASE CHAT HISTORY", "Owner-isolated conversations persist across browser sessions and devices.", BLUE),
        ("RETENTION GOVERNANCE", "System policy: 30/90/180/365 days or no expiry. Teams may choose a shorter period.", CYAN),
        ("COMPLETE DELETION", "Manual or scheduled cleanup removes messages, citations, artifacts, attachments, and metadata.", ORANGE),
        ("PERSISTENT INGESTION", "Mongo-backed jobs track extraction, chunking, embedding, indexing, failure, and retry.", GREEN),
        ("RESTART RECOVERY", "Worker leases and recovery scans resume interrupted document and attachment processing.", BLUE),
        ("CONSISTENT BACKUP", "A versioned local archive captures MongoDB, MinIO, and Qdrant with checksums.", ORANGE),
    ]
    for i, (title, body, color) in enumerate(feature_cards):
        col, row = i % 3, i // 3
        add_card(slide, title, body, Inches(0.68 + col * 4.18), Inches(1.92 + row * 2.05),
                 Inches(3.8), Inches(1.65), color, WHITE, f"{i + 1:02d}")
    add_rect(slide, Inches(0.72), Inches(6.12), Inches(11.85), Inches(0.58), PALE_CYAN, True, CYAN)
    add_text(slide, "Deletion audit stores who, when, reason, and counts—never deleted prompts, responses, citations, or diagram XML.",
             Inches(0.95), Inches(6.25), Inches(11.35), Inches(0.3), 12, INK, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # Slides were appended for stable python-pptx creation. Move them before the
    # original closing slide so they remain part of the main sales narrative.
    slide_id_list = prs.slides._sldIdLst
    for slide_id in feature_slide_ids:
        slide_id_list.remove(slide_id)
    for offset, slide_id in enumerate(feature_slide_ids):
        slide_id_list.insert(14 + offset, slide_id)

    # The original hardware backup slides now follow the new closing slide.
    for slide_number in range(20, 23):
        slide = prs.slides[slide_number - 1]
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame") or shape.left < Inches(12) or shape.top < Inches(6.8):
                continue
            runs = shape.text_frame.paragraphs[0].runs
            if runs:
                runs[0].text = f"{slide_number:02d}"
            else:
                shape.text_frame.paragraphs[0].text = f"{slide_number:02d}"
            break


def translate_deck():
    build_deck()
    prs = Presentation(TH_OUTPUT)
    missing = set()
    for slide in prs.slides:
        for shape in slide.shapes:
            if not hasattr(shape, "text_frame") or not shape.text.strip():
                continue
            source = normalize(shape.text)
            if source in TRANSLATIONS:
                shape.text_frame.paragraphs[0].text = TRANSLATIONS[source].replace("\n", "\v")
            elif any(ord(char) > 127 and char not in "•→—“”×–" for char in source):
                missing.add(source)
    if missing:
        raise RuntimeError("Missing English translations: " + repr(sorted(missing)))
    add_feature_slides(prs)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    translate_deck()
