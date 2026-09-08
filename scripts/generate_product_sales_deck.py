from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "presentations" / "InsightSphere_Product_Sales_Deck_TH.pptx"
LOGO = ROOT / "app" / "frontend" / "public" / "insight_sphere_new.png"

NAVY = RGBColor(11, 24, 39)
NAVY_2 = RGBColor(18, 36, 58)
BLUE = RGBColor(35, 126, 218)
CYAN = RGBColor(39, 187, 222)
ORANGE = RGBColor(255, 130, 24)
WHITE = RGBColor(255, 255, 255)
INK = RGBColor(31, 42, 55)
MUTED = RGBColor(96, 112, 128)
LIGHT = RGBColor(244, 248, 252)
PALE_BLUE = RGBColor(232, 243, 255)
PALE_CYAN = RGBColor(230, 250, 252)
PALE_ORANGE = RGBColor(255, 244, 229)
GREEN = RGBColor(28, 142, 91)
RED = RGBColor(195, 55, 55)

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
FONT = "Noto Sans Thai"


def set_bg(slide, color=WHITE):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_rect(slide, x, y, w, h, fill, radius=True, line=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, x, y, w, h)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = line or fill
    return shape


def add_text(slide, text, x, y, w, h, size=20, color=INK, bold=False,
             align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, font=FONT):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    paragraph.alignment = align
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = color
    return box


def add_rich_text(slide, parts, x, y, w, h, size=20, align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(x, y, w, h)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    for index, (text, color, bold) in enumerate(parts):
        run = paragraph.add_run() if index else paragraph.runs[0]
        run.text = text
        run.font.name = FONT
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
    return box


def add_header(slide, title, section=None, page=None, dark=False):
    fg = WHITE if dark else INK
    muted = RGBColor(178, 198, 218) if dark else MUTED
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(0.55), Inches(0.25), width=Inches(2.25))
    if section:
        add_text(slide, section.upper(), Inches(10.35), Inches(0.31), Inches(2.25), Inches(0.3),
                 10, muted, True, PP_ALIGN.RIGHT)
    add_text(slide, title, Inches(0.62), Inches(0.95), Inches(11.9), Inches(0.62), 27, fg, True)
    add_rect(slide, Inches(0.62), Inches(1.58), Inches(0.75), Inches(0.06), BLUE, False)
    if page is not None:
        add_text(slide, f"{page:02d}", Inches(12.18), Inches(7.05), Inches(0.5), Inches(0.25),
                 9, muted, True, PP_ALIGN.RIGHT)


def add_footer(slide, text="InsightSphere | Enterprise AI Agent & Knowledge Platform", dark=False):
    color = RGBColor(150, 174, 198) if dark else RGBColor(135, 148, 160)
    add_text(slide, text, Inches(0.62), Inches(7.02), Inches(8.5), Inches(0.25), 8, color)


def add_bullets(slide, items, x, y, w, h, size=18, color=INK, bullet_color=BLUE,
                spacing=10):
    row_h = h / max(len(items), 1)
    for i, item in enumerate(items):
        yy = y + row_h * i
        add_rect(slide, x, yy + Inches(0.12), Inches(0.11), Inches(0.11), bullet_color, False)
        add_text(slide, item, x + Inches(0.25), yy, w - Inches(0.25), row_h,
                 size, color, False, valign=MSO_ANCHOR.TOP)


def add_card(slide, title, body, x, y, w, h, accent=BLUE, fill=WHITE, badge=None):
    card = add_rect(slide, x, y, w, h, fill, True, RGBColor(220, 230, 240))
    card.line.width = Pt(1)
    add_rect(slide, x, y, Inches(0.09), h, accent, False)
    if badge:
        add_rect(slide, x + Inches(0.28), y + Inches(0.25), Inches(0.62), Inches(0.4), accent, True)
        add_text(slide, badge, x + Inches(0.28), y + Inches(0.25), Inches(0.62), Inches(0.4),
                 10, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        title_x = x + Inches(1.04)
        title_w = w - Inches(1.25)
    else:
        title_x = x + Inches(0.32)
        title_w = w - Inches(0.6)
    add_text(slide, title, title_x, y + Inches(0.22), title_w, Inches(0.48), 18, INK, True)
    add_text(slide, body, x + Inches(0.32), y + Inches(0.82), w - Inches(0.62), h - Inches(1.0),
             13, MUTED)
    return card


def add_metric(slide, label, value, x, y, w, accent=BLUE):
    add_rect(slide, x, y, w, Inches(1.05), WHITE, True, RGBColor(220, 230, 240))
    add_text(slide, value, x + Inches(0.2), y + Inches(0.13), w - Inches(0.4), Inches(0.42),
             20, accent, True, PP_ALIGN.CENTER)
    add_text(slide, label, x + Inches(0.15), y + Inches(0.58), w - Inches(0.3), Inches(0.3),
             10, MUTED, False, PP_ALIGN.CENTER)


def new_slide(prs, bg=WHITE):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_bg(slide, bg)
    return slide


def build_deck():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    # 1 — Cover
    slide = new_slide(prs, NAVY)
    add_rect(slide, Inches(0), Inches(0), Inches(4.35), SLIDE_H, NAVY_2, False)
    add_rect(slide, Inches(0), Inches(0), Inches(0.14), SLIDE_H, BLUE, False)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(0.65), Inches(0.65), width=Inches(3.1))
    add_text(slide, "Enterprise AI Agent\n& Knowledge Platform", Inches(4.95), Inches(1.15),
             Inches(7.45), Inches(1.65), 34, WHITE, True)
    add_text(slide, "เปลี่ยนความรู้และระบบงานขององค์กร\nให้เป็น AI ที่ค้นหา วิเคราะห์ และลงมือทำได้",
             Inches(4.98), Inches(3.05), Inches(6.95), Inches(1.15), 21,
             RGBColor(202, 219, 236))
    add_rect(slide, Inches(4.98), Inches(4.72), Inches(2.25), Inches(0.55), BLUE, True)
    add_text(slide, "PRODUCT OVERVIEW", Inches(4.98), Inches(4.72), Inches(2.25), Inches(0.55),
             11, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, "Private • Governed • Flexible Deployment", Inches(4.98), Inches(5.55),
             Inches(5.7), Inches(0.4), 13, RGBColor(140, 183, 225), True)
    add_text(slide, "Sales Presentation", Inches(4.98), Inches(6.65), Inches(2.2), Inches(0.3),
             10, RGBColor(144, 164, 184))

    # 2 — Executive proposition
    slide = new_slide(prs)
    add_header(slide, "AI Platform เดียว สำหรับความรู้ คน และระบบงาน", "Executive Summary", 2)
    add_text(slide, "InsightSphere ช่วยให้องค์กรสร้าง AI Agent ที่ตอบจากข้อมูลจริง เชื่อมต่อระบบเดิม\nและควบคุมสิทธิ์ได้ โดยเลือกใช้ Cloud, Local หรือ Hybrid ได้ตามนโยบาย",
             Inches(0.65), Inches(1.9), Inches(11.7), Inches(0.85), 21, INK, True)
    cards = [
        ("KNOW", "ค้นหาและอ้างอิงความรู้", "RAG พร้อม Citation ระดับไฟล์/หน้า และขอบเขตข้อมูล Personal, Team, Global", BLUE),
        ("THINK", "วิเคราะห์หลายรูปแบบ", "ข้อความ เอกสาร ภาพ OCR และ Evidence Matrix สำหรับงานตรวจข้อกำหนด", CYAN),
        ("ACT", "Agent เชื่อมต่อและลงมือทำ", "เรียก Tool, MCP, Database, API, Mail และ Messaging ภายใต้นโยบายองค์กร", ORANGE),
    ]
    for i, (badge, title, body, color) in enumerate(cards):
        add_card(slide, title, body, Inches(0.68 + i * 4.18), Inches(3.05), Inches(3.8), Inches(2.35), color, LIGHT, badge)
    add_text(slide, "ผลลัพธ์: จาก “ทดลองใช้ AI” สู่ “ระบบงาน AI ที่บริหารจัดการได้”",
             Inches(0.68), Inches(5.9), Inches(11.95), Inches(0.55), 20, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 3 — Pain points
    slide = new_slide(prs, LIGHT)
    add_header(slide, "ความท้าทายที่องค์กรพบเมื่อเริ่มใช้ Generative AI", "Customer Challenges", 3)
    pains = [
        ("ข้อมูลกระจัดกระจาย", "เอกสารอยู่หลายทีม หลายระบบ และค้นหาคำตอบที่เชื่อถือได้ยาก"),
        ("คำตอบไม่มีหลักฐาน", "AI ตอบได้ แต่ผู้ใช้ตรวจสอบแหล่งข้อมูลและหน้าที่อ้างอิงไม่ได้"),
        ("สิทธิ์เข้าถึงซับซ้อน", "ข้อมูลแต่ละทีมมีระดับความลับต่างกัน และต้องป้องกันการเห็นข้อมูลข้ามทีม"),
        ("เชื่อมต่อระบบยาก", "AI ที่ตอบอย่างเดียวไม่ลดขั้นตอน หากยังเรียก API, Database หรือ Tool ไม่ได้"),
        ("Vendor Lock-in", "ข้อกำหนดด้านต้นทุน ข้อมูล และโมเดลเปลี่ยนเร็ว แต่ระบบผูกกับผู้ให้บริการรายเดียว"),
        ("Production Governance", "ขาด Audit, Policy, Usage และวิธีบริหาร Agent/Model แบบศูนย์กลาง"),
    ]
    for i, (title, body) in enumerate(pains):
        col, row = i % 2, i // 2
        add_card(slide, title, body, Inches(0.68 + col * 6.15), Inches(1.9 + row * 1.55),
                 Inches(5.75), Inches(1.23), RED if i in {1, 2, 5} else ORANGE, WHITE, f"{i+1:02d}")
    add_footer(slide)

    # 4 — Platform overview
    slide = new_slide(prs, NAVY)
    add_header(slide, "แพลตฟอร์มที่รวม Knowledge, Agent และ Governance", "Solution", 4, True)
    add_text(slide, "USER EXPERIENCE", Inches(0.75), Inches(1.95), Inches(1.65), Inches(0.35), 10, CYAN, True)
    add_rect(slide, Inches(2.45), Inches(1.78), Inches(9.95), Inches(0.75), RGBColor(29, 55, 83), True)
    add_text(slide, "Chat & Ask  •  File / Image / Clipboard  •  Citation Viewer  •  Admin Portal",
             Inches(2.65), Inches(1.94), Inches(9.55), Inches(0.4), 15, WHITE, True, PP_ALIGN.CENTER)
    layers = [
        ("AI AGENTS", "Per-Agent Model • Prompt • RAG Policy • Tool Policy • LangGraph", BLUE),
        ("KNOWLEDGE", "Document Processing • OCR Routing • Vector Search • Evidence Matrix", CYAN),
        ("INTEGRATIONS", "MCP • REST API • Database • Mail • Messaging • Sub-Agent", ORANGE),
        ("GOVERNANCE", "Local + Entra ID • Team/Data Group ACL • Audit • Usage • Secrets", GREEN),
    ]
    for i, (label, body, color) in enumerate(layers):
        y = 2.85 + i * 0.88
        add_rect(slide, Inches(0.75), Inches(y), Inches(1.65), Inches(0.62), color, True)
        add_text(slide, label, Inches(0.75), Inches(y), Inches(1.65), Inches(0.62), 10, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_rect(slide, Inches(2.45), Inches(y), Inches(9.95), Inches(0.62), RGBColor(24, 47, 72), True)
        add_text(slide, body, Inches(2.68), Inches(y + 0.11), Inches(9.5), Inches(0.35), 13,
                 RGBColor(215, 227, 239), False, PP_ALIGN.CENTER)
    add_text(slide, "DEPLOYMENT", Inches(0.75), Inches(6.55), Inches(1.65), Inches(0.3), 10, CYAN, True)
    add_text(slide, "Azure  |  Local Containers  |  Hybrid Providers", Inches(2.45), Inches(6.5),
             Inches(9.95), Inches(0.4), 15, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, dark=True)

    # 5 — Knowledge
    slide = new_slide(prs)
    add_header(slide, "Enterprise Knowledge ที่ตอบได้และตรวจสอบได้", "Knowledge & RAG", 5)
    add_card(slide, "รับข้อมูลหลายขอบเขต", "Personal สำหรับผู้ใช้, Team + Data Group สำหรับงานร่วมกัน และ Global สำหรับความรู้ส่วนกลาง",
             Inches(0.68), Inches(1.93), Inches(3.78), Inches(1.75), BLUE, PALE_BLUE)
    add_card(slide, "Document Processing & OCR", "Routing ตามชนิดไฟล์ ทีม และลำดับความสำคัญ พร้อม Local parser, Docling และผู้ให้บริการ OCR",
             Inches(4.78), Inches(1.93), Inches(3.78), Inches(1.75), CYAN, PALE_CYAN)
    add_card(slide, "Citation และ Page Evidence", "คำตอบเชื่อมกลับไปยังไฟล์และหน้าต้นทาง พร้อม authorization ก่อนเปิดเอกสารทุกครั้ง",
             Inches(8.88), Inches(1.93), Inches(3.78), Inches(1.75), ORANGE, PALE_ORANGE)
    add_text(slide, "วงจรข้อมูลที่ควบคุมได้", Inches(0.72), Inches(4.1), Inches(3.0), Inches(0.4), 18, INK, True)
    flow = [("UPLOAD", BLUE), ("EXTRACT / OCR", CYAN), ("EMBED + INDEX", GREEN), ("RETRIEVE", ORANGE), ("CITE", BLUE)]
    for i, (label, color) in enumerate(flow):
        x = 0.72 + i * 2.48
        add_rect(slide, Inches(x), Inches(4.72), Inches(1.9), Inches(0.72), color, True)
        add_text(slide, label, Inches(x), Inches(4.72), Inches(1.9), Inches(0.72), 11, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        if i < len(flow) - 1:
            add_text(slide, "→", Inches(x + 1.92), Inches(4.81), Inches(0.48), Inches(0.4), 20, MUTED, True, PP_ALIGN.CENTER)
    add_text(slide, "Benefit", Inches(0.72), Inches(5.95), Inches(0.8), Inches(0.35), 12, BLUE, True)
    add_text(slide, "ลดเวลาค้นหา • เพิ่มความน่าเชื่อถือ • รักษาสิทธิ์ข้อมูล • รองรับการตรวจสอบย้อนหลัง",
             Inches(1.55), Inches(5.9), Inches(10.7), Inches(0.45), 17, INK, True)
    add_footer(slide)

    # 6 — Agents
    slide = new_slide(prs, LIGHT)
    add_header(slide, "AI Agent ที่ไม่ได้แค่ตอบ แต่ทำงานร่วมกับระบบองค์กร", "AI Agents", 6)
    add_text(slide, "กำหนด Agent แยกตามบทบาท พร้อม Model, Prompt, Data Scope และ Tool ที่เหมาะกับงาน",
             Inches(0.68), Inches(1.88), Inches(12), Inches(0.55), 19, INK, True)
    agents = [
        ("HR Agent", "ตอบนโยบายพนักงาน\nจากเอกสารที่ได้รับสิทธิ์", BLUE),
        ("Presales Agent", "ค้นข้อมูลผลิตภัณฑ์\nและสร้างคำตอบเชิงหลักฐาน", CYAN),
        ("Operations Agent", "ตรวจ Monitor / Log\nผ่าน MCP และสรุปเหตุการณ์", ORANGE),
        ("Data Agent", "Query ฐานข้อมูลแบบ read-only\nและสร้าง Chart", GREEN),
    ]
    for i, (title, body, color) in enumerate(agents):
        x = 0.68 + i * 3.08
        add_rect(slide, Inches(x), Inches(2.82), Inches(2.73), Inches(2.1), WHITE, True, RGBColor(216, 227, 238))
        add_rect(slide, Inches(x + 0.2), Inches(3.03), Inches(0.5), Inches(0.5), color, True)
        add_text(slide, "AI", Inches(x + 0.2), Inches(3.03), Inches(0.5), Inches(0.5), 11, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, title, Inches(x + 0.82), Inches(3.02), Inches(1.7), Inches(0.45), 16, INK, True)
        add_text(slide, body, Inches(x + 0.25), Inches(3.78), Inches(2.23), Inches(0.8), 13, MUTED, False, PP_ALIGN.CENTER)
    add_bullets(slide, [
        "Tool policy จำกัดจำนวนครั้ง เวลา และแหล่งที่เรียกใช้",
        "Sub-agent delegation พร้อมป้องกัน cycle และ depth เกินกำหนด",
        "Audit event เชื่อม model, tool และ trace โดยไม่เก็บ raw secrets",
    ], Inches(0.72), Inches(5.35), Inches(11.8), Inches(1.2), 14, INK, BLUE)
    add_footer(slide)

    # 7 — Document intelligence
    slide = new_slide(prs)
    add_header(slide, "จากเอกสารข้อกำหนด สู่ Evidence Matrix ที่ตรวจสอบได้", "Document Intelligence", 7)
    add_rect(slide, Inches(0.72), Inches(1.93), Inches(3.15), Inches(3.95), NAVY_2, True)
    add_text(slide, "INPUT", Inches(1.02), Inches(2.25), Inches(0.8), Inches(0.3), 11, CYAN, True)
    add_text(slide, "ข้อกำหนดงาน\nTOR / RFP / Policy\nChecklist / Standard", Inches(1.02), Inches(2.75), Inches(2.45), Inches(1.25), 21, WHITE, True)
    add_text(slide, "รองรับเอกสารแนบชั่วคราว\nหรือ Knowledge ที่องค์กรจัดเก็บ",
             Inches(1.02), Inches(4.55), Inches(2.35), Inches(0.7), 12, RGBColor(190, 210, 230))
    add_text(slide, "→", Inches(4.03), Inches(3.5), Inches(0.55), Inches(0.6), 27, BLUE, True, PP_ALIGN.CENTER)
    add_rect(slide, Inches(4.65), Inches(1.93), Inches(7.95), Inches(3.95), LIGHT, True, RGBColor(210, 223, 236))
    headers = ["Requirement", "Result", "Evidence", "Source", "Confidence / Gap"]
    widths = [2.05, 1.05, 2.05, 1.25, 1.15]
    x = 4.83
    for h, w in zip(headers, widths):
        add_rect(slide, Inches(x), Inches(2.25), Inches(w), Inches(0.58), PALE_BLUE, False, RGBColor(195, 212, 230))
        add_text(slide, h, Inches(x + 0.05), Inches(2.35), Inches(w - 0.1), Inches(0.3), 10, INK, True, PP_ALIGN.CENTER)
        x += w
    rows = [
        ("ระบบต้องมี Audit", "พบ", "Agent audit event", "Guide.pdf", "High"),
        ("รองรับ Local", "พบ", "Compose runtime", "Arch.pdf", "High"),
        ("ข้อกำหนดเฉพาะ", "ไม่พบ", "—", "—", "Gap"),
    ]
    for r, row in enumerate(rows):
        x = 4.83
        fill = WHITE if r % 2 == 0 else RGBColor(247, 250, 253)
        for c, (value, w) in enumerate(zip(row, widths)):
            add_rect(slide, Inches(x), Inches(2.83 + r * 0.78), Inches(w), Inches(0.78), fill, False, RGBColor(218, 228, 238))
            color = RED if value == "Gap" or value == "ไม่พบ" else INK
            add_text(slide, value, Inches(x + 0.08), Inches(2.98 + r * 0.78), Inches(w - 0.16), Inches(0.4),
                     10, color, value in {"พบ", "ไม่พบ", "High", "Gap"}, PP_ALIGN.CENTER)
            x += w
    add_text(slide, "Benefit: ตรวจข้อกำหนดได้เป็นระบบ เห็นทั้งหลักฐานและช่องว่าง ลดความเสี่ยงจากคำตอบที่ไม่มีข้อมูลรองรับ",
             Inches(0.75), Inches(6.25), Inches(11.9), Inches(0.5), 16, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 8 — Governance
    slide = new_slide(prs, NAVY)
    add_header(slide, "Security & Governance เป็นส่วนหนึ่งของ Platform", "Trust", 8, True)
    governance = [
        ("IDENTITY", "Local User + Microsoft Entra ID\nรองรับ bootstrap super admin", BLUE),
        ("ACCESS", "Team, Data Group, Personal/Global ACL\nตรวจสิทธิ์ทั้งค้นหาและเปิด Citation", CYAN),
        ("SECRETS", "Fernet encryption สำหรับ provider key\nไม่ส่ง secret กลับ frontend", ORANGE),
        ("TOOLS", "Allowlist, timeout, call limit\nread-only database enforcement", GREEN),
        ("AUDIT", "Trace model/tool activity\nพร้อม metadata โดยไม่เก็บ raw prompt/output", BLUE),
        ("ISOLATION", "History และ attachment แยกตาม user\nconversation และ data scope", CYAN),
    ]
    for i, (title, body, color) in enumerate(governance):
        col, row = i % 3, i // 3
        x, y = 0.72 + col * 4.18, 1.93 + row * 2.05
        add_rect(slide, Inches(x), Inches(y), Inches(3.78), Inches(1.7), RGBColor(24, 47, 72), True, RGBColor(48, 78, 107))
        add_rect(slide, Inches(x + 0.25), Inches(y + 0.25), Inches(0.62), Inches(0.38), color, True)
        add_text(slide, title, Inches(x + 0.25), Inches(y + 0.25), Inches(0.62), Inches(0.38), 8, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, body, Inches(x + 0.25), Inches(y + 0.78), Inches(3.28), Inches(0.68), 12, RGBColor(218, 230, 241))
    add_text(slide, "Designed for enterprise control — ไม่ลดทอน Security Contract เพื่อความสะดวก",
             Inches(0.75), Inches(6.42), Inches(11.9), Inches(0.4), 15, WHITE, True, PP_ALIGN.CENTER)
    add_footer(slide, dark=True)

    # 9 — Model/provider flexibility
    slide = new_slide(prs)
    add_header(slide, "เลือก Model และ Provider ให้เหมาะกับแต่ละ Agent", "Model Flexibility", 9)
    add_text(slide, "ไม่บังคับให้ทุกงานใช้โมเดลเดียว — แยก Chat Model, Embedding และ Processing Provider ได้",
             Inches(0.72), Inches(1.88), Inches(11.85), Inches(0.5), 19, INK, True)
    provider_groups = [
        ("Cloud AI", ["Azure OpenAI", "OpenAI", "Gemini", "Anthropic", "OpenRouter"], BLUE, PALE_BLUE),
        ("Local / Private", ["Ollama", "vLLM", "LocalAI", "OpenAI-compatible"], GREEN, RGBColor(233, 248, 240)),
        ("Specialized", ["Groq", "Mistral", "DeepSeek", "Together"], ORANGE, PALE_ORANGE),
    ]
    for i, (title, items, color, fill) in enumerate(provider_groups):
        x = 0.72 + i * 4.18
        add_rect(slide, Inches(x), Inches(2.7), Inches(3.78), Inches(2.45), fill, True, RGBColor(216, 228, 240))
        add_text(slide, title, Inches(x + 0.25), Inches(2.98), Inches(3.28), Inches(0.4), 18, color, True, PP_ALIGN.CENTER)
        add_text(slide, "\n".join(items), Inches(x + 0.38), Inches(3.55), Inches(3.02), Inches(1.25), 13, INK, False, PP_ALIGN.CENTER)
    add_bullets(slide, [
        "Portal ทดสอบ Tool Calling และ Vision capability ก่อนเปิดใช้งาน",
        "กำหนด Model ต่อ Agent เพื่อสมดุล Cost, Latency, Privacy และ Quality",
        "เปลี่ยน provider ได้โดยรักษา Agent, ACL และ Knowledge architecture เดิม",
    ], Inches(0.75), Inches(5.55), Inches(11.75), Inches(1.15), 14, INK, BLUE)
    add_footer(slide)

    # 10 — Deployment
    slide = new_slide(prs, LIGHT)
    add_header(slide, "Deployment ที่ยืดหยุ่น: Azure, Local หรือ Hybrid", "Deployment", 10)
    modes = [
        ("AZURE", "Azure OpenAI\nAzure AI Search\nBlob / Cosmos-compatible\nEntra ID", BLUE),
        ("LOCAL", "Docker Compose\nMongoDB + MinIO\nQdrant + Ollama/BGE-M3\nLocal Authentication", GREEN),
        ("HYBRID", "Local data plane\nเลือก Cloud LLM\nหรือ Local Model\nตาม Agent / Policy", ORANGE),
    ]
    for i, (title, body, color) in enumerate(modes):
        x = 0.78 + i * 4.15
        add_rect(slide, Inches(x), Inches(2.0), Inches(3.65), Inches(3.4), WHITE, True, RGBColor(207, 221, 235))
        add_rect(slide, Inches(x), Inches(2.0), Inches(3.65), Inches(0.72), color, True)
        add_text(slide, title, Inches(x), Inches(2.0), Inches(3.65), Inches(0.72), 17, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, body, Inches(x + 0.38), Inches(3.1), Inches(2.89), Inches(1.75), 16, INK, False, PP_ALIGN.CENTER)
    add_text(slide, "แยก Provider Interface ชัดเจน — รักษา Azure behavior พร้อมเปิดทางสู่ Self-hosted",
             Inches(0.78), Inches(5.85), Inches(11.75), Inches(0.48), 17, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 11 — Use cases
    slide = new_slide(prs)
    add_header(slide, "ตัวอย่าง Use Case ที่สร้างมูลค่าได้เร็ว", "Use Cases", 11)
    use_cases = [
        ("Enterprise Q&A", "ตอบคำถามนโยบาย คู่มือ และความรู้ภายใน พร้อม Citation", BLUE),
        ("RFP / Compliance", "เปรียบเทียบข้อกำหนดกับหลักฐานใน Knowledge และแสดง Gap", ORANGE),
        ("IT Operations", "ค้น Monitor/Log ผ่าน MCP สรุปเหตุการณ์ และแนะนำขั้นตอนแก้ไข", CYAN),
        ("Data Assistant", "Query ฐานข้อมูลแบบควบคุมสิทธิ์ สรุปและสร้างกราฟ", GREEN),
        ("Presales Copilot", "ค้น Product Knowledge สร้างคำตอบที่ตรวจสอบแหล่งข้อมูลได้", BLUE),
        ("Workflow Agent", "เรียก API, Mail, Messaging และ Sub-agent ตาม Tool policy", ORANGE),
    ]
    for i, (title, body, color) in enumerate(use_cases):
        col, row = i % 3, i // 3
        add_card(slide, title, body, Inches(0.68 + col * 4.18), Inches(1.92 + row * 2.18),
                 Inches(3.8), Inches(1.78), color, LIGHT, f"{i+1:02d}")
    add_footer(slide)

    # 12 — Benefits
    slide = new_slide(prs, NAVY)
    add_header(slide, "Business Benefits ที่เชื่อมโยงกับการวัดผล", "Business Value", 12, True)
    benefits = [
        ("PRODUCTIVITY", "ลดเวลาค้นหา สรุป และจัดทำคำตอบซ้ำ", "วัด: เวลาเฉลี่ยต่อเคส", BLUE),
        ("QUALITY", "คำตอบมีหลักฐานและเห็นช่องว่างข้อมูล", "วัด: acceptance / rework rate", CYAN),
        ("CONTROL", "สิทธิ์ Agent Tool และข้อมูลบริหารจากศูนย์กลาง", "วัด: policy violation / audit coverage", GREEN),
        ("FLEXIBILITY", "เปลี่ยน Model และ Deployment ตามงานได้", "วัด: cost per task / latency", ORANGE),
    ]
    for i, (title, body, metric, color) in enumerate(benefits):
        y = 1.93 + i * 1.18
        add_rect(slide, Inches(0.75), Inches(y), Inches(2.15), Inches(0.78), color, True)
        add_text(slide, title, Inches(0.75), Inches(y), Inches(2.15), Inches(0.78), 11, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, body, Inches(3.2), Inches(y + 0.05), Inches(5.6), Inches(0.36), 15, WHITE, True)
        add_text(slide, metric, Inches(8.95), Inches(y + 0.07), Inches(3.6), Inches(0.32), 12, RGBColor(175, 204, 231), False, PP_ALIGN.RIGHT)
    add_text(slide, "เริ่มจาก Baseline ของลูกค้า แล้ววัดผล Pilot ก่อนสรุป ROI — ไม่ใช้ตัวเลขสมมติ",
             Inches(0.75), Inches(6.55), Inches(11.8), Inches(0.35), 14, CYAN, True, PP_ALIGN.CENTER)
    add_footer(slide, dark=True)

    # 13 — Delivery approach
    slide = new_slide(prs, LIGHT)
    add_header(slide, "แนวทางเริ่มใช้งานแบบลดความเสี่ยง", "Engagement", 13)
    phases = [
        ("01", "DISCOVER", "เลือก 1–2 Use Case\nกำหนด Data/ACL/KPI", BLUE),
        ("02", "PILOT", "เชื่อมข้อมูลจริง\nตั้ง Agent และ Guardrails", CYAN),
        ("03", "VALIDATE", "User Acceptance\nSecurity & KPI Review", ORANGE),
        ("04", "SCALE", "Production hardening\nOperations & Enablement", GREEN),
    ]
    for i, (num, title, body, color) in enumerate(phases):
        x = 0.72 + i * 3.1
        add_rect(slide, Inches(x), Inches(2.15), Inches(2.72), Inches(3.25), WHITE, True, RGBColor(213, 225, 237))
        add_rect(slide, Inches(x + 0.95), Inches(2.48), Inches(0.82), Inches(0.82), color, True)
        add_text(slide, num, Inches(x + 0.95), Inches(2.48), Inches(0.82), Inches(0.82), 18, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, title, Inches(x + 0.28), Inches(3.62), Inches(2.16), Inches(0.42), 15, color, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x + 0.3), Inches(4.25), Inches(2.12), Inches(0.72), 13, INK, False, PP_ALIGN.CENTER)
        if i < 3:
            add_text(slide, "→", Inches(x + 2.72), Inches(3.36), Inches(0.37), Inches(0.5), 20, MUTED, True, PP_ALIGN.CENTER)
    add_bullets(slide, [
        "Pilot ใช้ข้อมูลจริงในขอบเขตจำกัด พร้อมผู้รับผิดชอบข้อมูล",
        "Production scope รวม TLS, Backup/Restore, Monitoring, Capacity และ Support model",
    ], Inches(0.76), Inches(5.85), Inches(11.7), Inches(0.72), 13, INK, BLUE)
    add_footer(slide)

    # 14 — Differentiators
    slide = new_slide(prs)
    add_header(slide, "ทำไม InsightSphere เหมาะกับ Enterprise AI", "Differentiators", 14)
    differentiators = [
        ("Evidence-first", "Citation, Page, Evidence และ Gap ถูกออกแบบเป็น workflow ไม่ใช่ส่วนเสริม", BLUE),
        ("Access-aware", "Team/Data Group ACL อยู่ใน ingestion, retrieval และ citation authorization", GREEN),
        ("Agentic but governed", "Agent เรียก Tool ได้จริงภายใต้ policy, permission และ audit", ORANGE),
        ("Provider-neutral", "Cloud, Local และ Hybrid เปลี่ยนตามข้อกำหนดโดยไม่เขียน platform ใหม่", CYAN),
        ("Built for operations", "Admin Portal, model capability test, persistent jobs และ recovery", BLUE),
        ("Expandable architecture", "MCP, LangGraph, sub-agent และ provider adapter รองรับการเติบโต", GREEN),
    ]
    for i, (title, body, color) in enumerate(differentiators):
        col, row = i % 2, i // 2
        add_card(slide, title, body, Inches(0.72 + col * 6.15), Inches(1.93 + row * 1.52),
                 Inches(5.72), Inches(1.2), color, LIGHT, f"{i+1:02d}")
    add_text(slide, "One platform. Your data. Your model choice. Your governance.",
             Inches(0.72), Inches(6.5), Inches(11.85), Inches(0.4), 18, BLUE, True, PP_ALIGN.CENTER)
    add_footer(slide)

    # 15 — Close
    slide = new_slide(prs, NAVY)
    add_rect(slide, Inches(0), Inches(0), Inches(0.14), SLIDE_H, BLUE, False)
    if LOGO.exists():
        slide.shapes.add_picture(str(LOGO), Inches(0.75), Inches(0.62), width=Inches(3.0))
    add_text(slide, "เริ่มจาก Use Case ที่เห็นผล\nแล้วขยายเป็น Enterprise AI Platform",
             Inches(0.78), Inches(1.8), Inches(8.9), Inches(1.35), 31, WHITE, True)
    add_text(slide, "ข้อเสนอสำหรับ Next Step", Inches(0.82), Inches(3.55), Inches(3.2), Inches(0.4), 16, CYAN, True)
    add_bullets(slide, [
        "Workshop: Use Case, Data Scope, Security และ KPI",
        "Technical Assessment: Deployment และ Integration",
        "Pilot Plan: ขอบเขต ผู้ใช้ ระยะเวลา และเกณฑ์ความสำเร็จ",
    ], Inches(0.84), Inches(4.15), Inches(8.5), Inches(1.65), 17, WHITE, CYAN)
    add_rect(slide, Inches(9.85), Inches(2.05), Inches(2.3), Inches(2.3), BLUE, True)
    add_text(slide, "LET'S\nBUILD\nTRUSTED AI", Inches(9.85), Inches(2.05), Inches(2.3), Inches(2.3),
             20, WHITE, True, PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
    add_text(slide, "Thank you", Inches(0.82), Inches(6.55), Inches(2.2), Inches(0.4), 16, RGBColor(168, 194, 219), True)

    # 16 — Backup: capacity framing
    slide = new_slide(prs, LIGHT)
    add_header(slide, "การประเมิน Capacity ต้องเริ่มจาก Concurrent Workload", "Backup | Capacity", 16)
    add_text(slide, "จำนวน Account ไม่เท่ากับจำนวนงาน AI ที่ประมวลผลพร้อมกัน",
             Inches(0.72), Inches(1.87), Inches(11.85), Inches(0.48), 20, INK, True)
    capacity_metrics = [
        ("REGISTERED", "ผู้ใช้ทั้งหมด", "ใช้วางแผน Identity\nและข้อมูลระยะยาว", BLUE),
        ("ACTIVE", "ผู้ใช้ที่ใช้งานจริง", "ดู Daily/Monthly Active\nและรูปแบบการใช้งาน", CYAN),
        ("CONCURRENT", "คำขอที่กำลังทำงาน", "กำหนด App replicas,\nconnection และ queue", ORANGE),
        ("GENERATION", "LLM ที่สร้างคำตอบพร้อมกัน", "กำหนด GPU slots\nหรือ Cloud quota", GREEN),
    ]
    for i, (badge, title, body, color) in enumerate(capacity_metrics):
        x = 0.68 + i * 3.08
        add_rect(slide, Inches(x), Inches(2.72), Inches(2.73), Inches(2.22), WHITE, True, RGBColor(216, 227, 238))
        add_rect(slide, Inches(x + 0.24), Inches(2.96), Inches(1.1), Inches(0.38), color, True)
        add_text(slide, badge, Inches(x + 0.24), Inches(2.96), Inches(1.1), Inches(0.38), 8, WHITE, True,
                 PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
        add_text(slide, title, Inches(x + 0.24), Inches(3.55), Inches(2.25), Inches(0.45), 15, INK, True, PP_ALIGN.CENTER)
        add_text(slide, body, Inches(x + 0.24), Inches(4.1), Inches(2.25), Inches(0.62), 12, MUTED, False, PP_ALIGN.CENTER)
    add_rect(slide, Inches(0.72), Inches(5.42), Inches(11.85), Inches(0.92), PALE_BLUE, True, RGBColor(193, 215, 237))
    add_text(slide, "Capacity โดยประมาณ = Peak concurrent requests × เวลาประมวลผล × ขนาด Context / Tool workload",
             Inches(0.98), Inches(5.63), Inches(11.3), Inches(0.34), 16, BLUE, True, PP_ALIGN.CENTER)
    add_text(slide, "ตัวเลขในเอกสารนี้เป็น Planning baseline — ต้องยืนยันด้วย Load Test จาก use case และ model จริง",
             Inches(0.8), Inches(6.48), Inches(11.7), Inches(0.32), 11, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide)

    # 17 — Backup: hardware sizing
    slide = new_slide(prs)
    add_header(slide, "Hardware Sizing Baseline สำหรับ Local / Hybrid", "Backup | Sizing", 17)
    headers = ["ระดับ", "Compute / Memory", "Storage / GPU", "แนวทางรองรับเบื้องต้น"]
    widths = [1.75, 3.1, 3.25, 4.1]
    x = 0.62
    for header, width in zip(headers, widths):
        add_rect(slide, Inches(x), Inches(1.82), Inches(width), Inches(0.58), NAVY_2, False, NAVY_2)
        add_text(slide, header, Inches(x + 0.08), Inches(1.94), Inches(width - 0.16), Inches(0.28), 11, WHITE, True, PP_ALIGN.CENTER)
        x += width
    sizing_rows = [
        ("Dev / Demo", "4–8 vCPU\nRAM 16 GB", "SSD 200 GB\nไม่บังคับ GPU", "1–5 ผู้ใช้พร้อมกัน\nเหมาะกับทดสอบและสาธิต"),
        ("Pilot\nCloud LLM", "8 vCPU\nRAM 32 GB", "NVMe 500 GB\nLLM ใช้ Cloud", "ประมาณ 10–30 concurrent Chat\nขึ้นกับ provider quota และ tools"),
        ("Pilot\nLocal LLM 8B", "8–16 vCPU\nRAM 32–64 GB", "NVMe 500 GB+\nGPU VRAM 16–24 GB", "1–2 local generations พร้อมกัน\nเหมาะกับ light / staggered usage"),
        ("Department / HA", "App 2+ nodes\nnode ละ 4–8 vCPU / 16 GB", "แยก Stateful services\nและ GPU inference node", "ประมาณ 30–100 concurrent users\nหลังผ่าน Load Test"),
    ]
    row_y = 2.4
    row_h = 0.98
    for row_index, row in enumerate(sizing_rows):
        x = 0.62
        fill = WHITE if row_index % 2 == 0 else LIGHT
        for col_index, (value, width) in enumerate(zip(row, widths)):
            add_rect(slide, Inches(x), Inches(row_y + row_index * row_h), Inches(width), Inches(row_h), fill, False,
                     RGBColor(211, 222, 233))
            color = BLUE if col_index == 0 else INK
            add_text(slide, value, Inches(x + 0.12), Inches(row_y + row_index * row_h + 0.13),
                     Inches(width - 0.24), Inches(row_h - 0.2), 11, color, col_index == 0,
                     PP_ALIGN.CENTER, MSO_ANCHOR.MIDDLE)
            x += width
    add_text(slide, "หมายเหตุ: พื้นที่จัดเก็บต้องคำนวณเพิ่มจากเอกสารต้นฉบับ + vector index + chat/attachment retention + backup",
             Inches(0.72), Inches(6.52), Inches(11.85), Inches(0.34), 11, MUTED, False, PP_ALIGN.CENTER)
    add_footer(slide)

    # 18 — Backup: scaling roadmap
    slide = new_slide(prs, NAVY)
    add_header(slide, "จาก Single-node Pilot สู่ Production Scale", "Backup | Scale", 18, True)
    add_text(slide, "CURRENT PILOT", Inches(0.72), Inches(1.9), Inches(2.0), Inches(0.35), 11, CYAN, True)
    add_rect(slide, Inches(0.72), Inches(2.35), Inches(4.25), Inches(3.55), RGBColor(24, 47, 72), True, RGBColor(48, 78, 107))
    add_bullets(slide, [
        "Compose เครื่องเดียว — ไม่มี automatic failover",
        "Hypercorn application process เดียว",
        "Ollama parallel = 1 เป็นคอขวด Local LLM",
        "MongoDB, MinIO และ Qdrant เป็น single instance",
        "Ingestion อยู่ใน App และจำกัด 3 งานต่อ process",
    ], Inches(1.0), Inches(2.68), Inches(3.7), Inches(2.85), 12, WHITE, ORANGE, 6)
    add_text(slide, "→", Inches(5.15), Inches(3.6), Inches(0.55), Inches(0.55), 28, CYAN, True, PP_ALIGN.CENTER)
    add_text(slide, "TARGET PRODUCTION", Inches(5.92), Inches(1.9), Inches(2.3), Inches(0.35), 11, CYAN, True)
    target_cards = [
        ("INGRESS + APP", "Load balancer และ App 2+ replicas", BLUE),
        ("ASYNC WORKERS", "Queue แยก OCR / ingestion / cleanup", ORANGE),
        ("SHARED DATA", "Mongo replica, Object storage, Qdrant replication", GREEN),
        ("INFERENCE", "Cloud quota หรือ GPU pool / vLLM", CYAN),
    ]
    for i, (title, body, color) in enumerate(target_cards):
        col, row = i % 2, i // 2
        x, y = 5.92 + col * 3.25, 2.35 + row * 1.78
        add_rect(slide, Inches(x), Inches(y), Inches(2.95), Inches(1.48), RGBColor(24, 47, 72), True, RGBColor(48, 78, 107))
        add_rect(slide, Inches(x + 0.2), Inches(y + 0.2), Inches(0.1), Inches(1.08), color, False)
        add_text(slide, title, Inches(x + 0.45), Inches(y + 0.2), Inches(2.25), Inches(0.35), 12, WHITE, True)
        add_text(slide, body, Inches(x + 0.45), Inches(y + 0.68), Inches(2.25), Inches(0.5), 11, RGBColor(197, 216, 234))
    add_text(slide, "Scale trigger: p95 latency, queue depth, provider throttling, GPU utilization และ ingestion backlog",
             Inches(0.76), Inches(6.42), Inches(11.8), Inches(0.35), 13, CYAN, True, PP_ALIGN.CENTER)
    add_footer(slide, dark=True)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_deck()
