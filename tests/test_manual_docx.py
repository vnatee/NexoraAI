from pathlib import Path
from xml.etree import ElementTree
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
WORD_NAMESPACE = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
MANUALS = {
    "NexoraAI-User-Manual.docx": 8,
    "NexoraAI-Administrator-Manual.docx": 11,
}


def test_word_manuals_are_readable_and_have_expected_placeholders():
    for filename, expected_placeholders in MANUALS.items():
        path = ROOT / "docs" / filename
        document = Document(path)
        placeholders = sum(
            "SCREENSHOT PLACEHOLDER" in paragraph.text
            for table in document.tables
            for row in table.rows
            for cell in row.cells
            for paragraph in cell.paragraphs
        )

        assert placeholders == expected_placeholders
        assert len(document.paragraphs) > 100
        assert len(document.tables) > expected_placeholders


def test_word_fields_are_nested_in_runs_for_microsoft_word_compatibility():
    for filename in MANUALS:
        path = ROOT / "docs" / filename
        field_count = 0
        with ZipFile(path) as archive:
            bad_direct_children = []
            for member in archive.namelist():
                if not member.startswith("word/") or not member.endswith(".xml"):
                    continue
                root = ElementTree.fromstring(archive.read(member))
                for paragraph in root.iter(f"{WORD_NAMESPACE}p"):
                    for child in paragraph:
                        if child.tag in {f"{WORD_NAMESPACE}fldChar", f"{WORD_NAMESPACE}instrText"}:
                            bad_direct_children.append((member, child.tag))
                for run in root.iter(f"{WORD_NAMESPACE}r"):
                    field_count += sum(child.tag == f"{WORD_NAMESPACE}fldChar" for child in run)

        assert bad_direct_children == []
        assert field_count == 9
